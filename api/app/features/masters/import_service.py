"""The bulk product importer: a spreadsheet in, a catalogue out.

The fourth member of the import family — `geo.import_geography`,
`masters.import_serials` and `tickets.import_service` are the other three —
and it keeps their contract exactly:

  * **Two passes over one file.** A dry run writes nothing and returns exactly
    what the commit would do, so every number the console shows is the
    server's own count rather than a guess made in a browser that never parsed
    the file.
  * **Rejected rows never block the good ones.** A bad cell costs that row,
    not the upload. The count is exact; the list is capped.
  * **Nothing is parsed in the client.** No spreadsheet library ships to a
    browser.

What is different here is the UNIT of work. The ticket importer's rows name
tickets, and a product the file does not recognise is a side-effect it has to
submit on the fly. Here a row IS a product — the Category column names the
chain it files under (the same chain grammar `tickets.import_service` reads,
duplicated rather than imported: hard rule 4 forbids one slice importing
another's service, and `_iter_rows` / `_cell_text` already make exactly this
trade three times over), and the rest of the row is what `ModelCreateRequest`
/ `ProductSubmitRequest` already ask a single save for.

## Create, or update — matched on (node, brand, name)

Re-uploading the same sheet is the normal way to use this: a vendor fixes a
capacity typo and uploads again, or a price list changes for forty products at
once. A row whose chain, brand and name already match a live product UPDATES
it; anything else CREATES one. There is no separate "Reference" column the
way the ticket importer needs one — `tickets.external_ref` exists because a
ticket has no other handle to repeat-detect on (`serial_number` is
deliberately not unique), but a product's own uniqueness constraint
(`uq_product_models_node_brand_name_lower`) already is that handle.

## Approval semantics mirror `create_model` vs `submit_model` exactly

A STAFF import creates a product APPROVED and PRICED, the same as
`create_model` — both prices are required on every row that CREATES one. A
VENDOR import creates it PENDING and UNPRICED, the same as `submit_model` —
the two price columns are not even read for a vendor's upload, because a
vendor has no field to send a price in anywhere else in this codebase either.

**A bulk edit must never reach around the approval rank floor.** A sheet that
matches an EXISTING product staff did not create directly can only update one
that is already `approved` — a `pending` or `rejected` row is somebody else's
decision (`masters.approve` + a National-Head floor), and `masters.edit` (what
this importer runs under for staff) must not let a lower rank price or approve
a product by naming it in a spreadsheet instead of on the Approvals screen.
A vendor's own update, by contrast, is `update_own_model`'s bulk twin: an
`approved` row keeps its prices and stays approved, and a `rejected` one that
the sheet still names is read as the vendor answering the refusal and goes
back to `pending` — exactly as editing it by hand would.
"""

import csv
import dataclasses
import datetime
import hashlib
import io
import json
import re
import uuid
from collections.abc import Iterator

from fastapi import HTTPException, status
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import Principal
from app.core.errors import AppError
from app.core.notifications import notify
from app.core.product_tree import (
    APPROVED,
    CERTIFY_DEPTH,
    MAX_PARAMETERS,
    PENDING,
    REJECTED,
)
from app.core.realtime import publish_notification
from app.core.service_types import DEFAULT_SERVICE_TYPES, SERVICE_TYPES
from app.features.masters.schemas import (
    CategoryToCreateOut,
    ProductImportReject,
    ProductImportReport,
)
from app.models.product import ProductModel
from app.models.vendor import Vendor
from app.models.vendor_brand import VendorBrand

#: A big catalogue upload is still a handful of round trips, not one per row —
#: see the batched reads below — so this is generous compared to the ticket
#: importer's 500. There is no per-row WhatsApp or credit charge here, just a
#: write.
MAX_PRODUCT_IMPORT_ROWS = 1000

#: Generous for a thousand rows of text; the same read-with-a-ceiling guard
#: every sibling importer uses — the declared size is not trusted.
MAX_PRODUCT_UPLOAD_BYTES = 8 * 1024 * 1024

#: How many rejects travel back. The COUNT is always exact; this caps the list
#: so a catastrophically wrong file cannot return a huge response.
MAX_PRODUCT_REJECTS_RETURNED = 200

_SHEET_NAME = "Products"
_README_SHEET = "How to fill this in"

_MARKER_RE = re.compile(r"\(([^()]*)\)\s*$")


# ── the template ─────────────────────────────────────────────────────────────


@dataclasses.dataclass(frozen=True)
class _Column:
    key: str
    header: str
    #: Everything that normalises to one of these maps to this column, so
    #: `Product Name`, `product name` and `MODEL` all land.
    aliases: tuple[str, ...]
    width: int
    example_a: str = ""
    example_b: str = ""


def _norm(raw: object) -> str:
    """Lowercased, letters and digits only — the header matcher every sibling uses."""
    return re.sub(r"[^a-z0-9]", "", str(raw or "").lower())


COLUMNS: tuple[_Column, ...] = (
    _Column(
        "category",
        "Category",
        ("category", "categories", "categorychain", "categorypath"),
        52,
        "Electronics, Television, Android TV(islastsubcategory)",
        "Home Appliances, Washing Machine(islastsubcategory)",
    ),
    _Column("brand", "Brand", ("brand", "brandname"), 16, "Meridian", "Sunview"),
    _Column(
        "name",
        "Product Name",
        ("productname", "modelname", "product", "model", "name"),
        26,
        "43 inch 4K UHD",
        "7 kg Front Load",
    ),
    _Column(
        "serviceTypes",
        "Service Types",
        ("servicetypes", "servicetype", "type", "requesttype"),
        30,
        "Installation + Demo",
        "Installation + Demo, Service",
    ),
    _Column("capacity", "Capacity", ("capacity", "size", "rating"), 16, "43 inch", "7 kg"),
    _Column(
        "warrantyMonths",
        "Warranty (months)",
        ("warrantymonths", "warranty"),
        16,
        "12",
        "24",
    ),
    _Column(
        "technicianPayoutPaise",
        "Technician Payout (paise)",
        ("technicianpayoutpaise", "technicianpayout", "payoutpaise", "payout"),
        24,
        "35000",
        "45000",
    ),
    _Column(
        "vendorPricePaise",
        "Vendor Price (paise)",
        ("vendorpricepaise", "vendorprice", "pricepaise", "price"),
        20,
        "180000",
        "220000",
    ),
    _Column(
        "parameters",
        "Parameters",
        ("parameters", "specs", "fields"),
        40,
        "RAM:8 GB; Storage:128 GB",
        "",
    ),
)

#: The columns a file must carry. Everything else may be blank or absent.
REQUIRED_COLUMNS = ("category", "name")

_BY_KEY = {c.key: c for c in COLUMNS}


def build_product_template() -> io.BytesIO:
    """The starter file: one sheet of columns, one sheet of rules."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = _SHEET_NAME
    sheet.append([c.header for c in COLUMNS])
    sheet.append([c.example_a for c in COLUMNS])
    sheet.append([c.example_b for c in COLUMNS])
    for i, col in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[
            openpyxl.utils.get_column_letter(i)
        ].width = col.width

    notes = book.create_sheet(_README_SHEET)
    for line in _README_LINES:
        notes.append([line])
    notes.column_dimensions["A"].width = 110

    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


_README_LINES: tuple[str, ...] = (
    "How to fill this in",
    "",
    "Every row is one product. Fill in the Products sheet and upload it.",
    "Uploading the same sheet again UPDATES the products it already created —",
    "matched by category, brand and name — instead of adding them twice.",
    "Rows that cannot be read are listed with a reason and do not stop the rest.",
    "",
    "Category",
    "  Name the whole chain, from the top-level category down to the level the",
    "  product is filed under, separated by commas:",
    "",
    "      Electronics, Television, Android TV(islastsubcategory)",
    "",
    "  Mark the LAST name with (islastsubcategory). Anything in the chain that",
    "  does not exist yet is created — you are shown exactly what before",
    "  anything is saved. A chain can be up to 6 names long.",
    "",
    "  If a category name contains a comma of its own, separate the names with >",
    "  instead. Then every comma in the cell is part of a name.",
    "",
    "Brand",
    "  One of your approved brands. Leave it blank only if you have exactly one.",
    "",
    "Service Types",
    "  One or more of: Installation + Demo, Tech Visit, Service — separated by",
    "  commas. Leave it blank for Installation + Demo.",
    "",
    "Warranty (months)",
    "  A whole number from 0 to 240. Leave it blank if you do not know it yet.",
    "",
    "Technician Payout (paise) and Vendor Price (paise)",
    "  Staff only — what the job is worth to each side, in paise (₹1 = 100).",
    "  A vendor's own upload ignores these two columns: a vendor never sets",
    "  either figure, and a National Head prices the product on Approvals.",
    "  A new product needs both figures; leave both blank on a row that is",
    "  only updating an existing product's other details.",
    "",
    "Parameters",
    "  Free-text specs as Name:Value, separated by semicolons:",
    "",
    "      RAM:8 GB; Storage:128 GB",
)


# ── reading the file ─────────────────────────────────────────────────────────


def _bad_request(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=detail)


def _iter_rows(data: bytes, filename: str) -> Iterator[tuple]:
    """Yield raw tuples from .xlsx or .csv, streaming in both cases.

    A near-twin of `tickets.import_service._iter_rows`, kept here for the
    reason that one gives: hard rule 4 forbids one slice importing another's
    service, and sharing fifteen lines is not worth coupling the product
    master to tickets.

    Prefers the sheet named `Products`, the way the ticket importer prefers
    `Tickets` — this template has two sheets, and a reordered tab must not be
    parsed as data.
    """
    if filename.lower().endswith(".csv"):
        decoded = data.decode("utf-8-sig", errors="replace")
        yield from csv.reader(io.StringIO(decoded))
        return

    try:
        import openpyxl
    except ModuleNotFoundError:  # pragma: no cover - dependency is in requirements
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Spreadsheet support is not installed on this server",
        )
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise _bad_request("That file could not be opened as a spreadsheet")
    try:
        name = next(
            (s for s in book.sheetnames if _norm(s) == _norm(_SHEET_NAME)),
            book.sheetnames[0],
        )
        yield from book[name].iter_rows(values_only=True)
    finally:
        book.close()


def _cell_text(value: object) -> str:
    """One spreadsheet cell as the string a human meant by it.

    The float branch matters here exactly as it does in every sibling: a
    price or a warranty typed as a plain number arrives from openpyxl as a
    FLOAT, and `str()` of that is `"35000.0"` rather than `"35000"`.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# ── the Category chain — duplicated from `tickets.import_service` ───────────
#
# Hard rule 4: slices never import each other. `tickets.import_service` is the
# one other place this grammar is read, and it is off-limits to touch for this
# change, so the parser is copied rather than shared — the same trade this
# module's docstring already names for `_iter_rows` / `_cell_text`.


def _is_marker(token: str) -> bool:
    return _norm(token) == "islastsubcategory"


def _looks_like_marker(token: str) -> bool:
    normalised = _norm(token)
    return not _is_marker(normalised) and "last" in normalised and "sub" in normalised


@dataclasses.dataclass
class _Chain:
    """One Category cell, parsed."""

    segments: list[str]
    code: str | None = None
    reason: str | None = None

    @property
    def key(self) -> tuple[str, ...]:
        return tuple(s.lower() for s in self.segments)


def parse_category_chain(raw: str) -> _Chain:
    """Split a Category cell into names and check the marker.

    Everything answerable from the cell alone. Depth, leaf conflicts and
    anything needing a row from the database are `_resolve_chains`'s job.
    """
    cell = (raw or "").strip()
    if not cell:
        return _Chain(
            [],
            "CATEGORY_EMPTY",
            "No category. Name the chain, ending with the last sub-category.",
        )

    delimiter = ">" if ">" in cell else ","
    raw_segments = cell.split(delimiter)

    if raw_segments and not raw_segments[0].strip():
        return _Chain(
            [],
            "CATEGORY_LEADING_DELIMITER",
            f'The chain starts with "{delimiter}". Remove it, or add the '
            "category before it.",
        )
    if raw_segments and not raw_segments[-1].strip():
        return _Chain(
            [],
            "CATEGORY_TRAILING_DELIMITER",
            f'The chain ends with "{delimiter}". Remove it, or add the missing '
            "sub-category after it.",
        )

    names: list[str] = []
    marked: list[int] = []
    for index, piece in enumerate(raw_segments):
        segment = re.sub(r"\s+", " ", piece.strip())
        found = _MARKER_RE.search(segment)
        if found is not None:
            inner = found.group(1)
            if _is_marker(inner):
                marked.append(index)
                segment = segment[: found.start()].strip()
            elif _looks_like_marker(inner):
                return _Chain(
                    [],
                    "CATEGORY_MARKER_MISSPELT",
                    f'"({inner})" is not the marker. Write it exactly as '
                    "(islastsubcategory).",
                )
        if not segment:
            if index in marked:
                return _Chain(
                    [],
                    "CATEGORY_MARKER_NO_NAME",
                    "The marker has no name in front of it. Write it on the "
                    'last sub-category, like "Android TV(islastsubcategory)".',
                )
            neighbours = [n for n in (names[-1] if names else None,) if n]
            between = f' after "{neighbours[0]}"' if neighbours else ""
            return _Chain(
                [],
                "CATEGORY_EMPTY_LEVEL",
                f"There is an empty name{between}. Remove the extra "
                f'"{delimiter}".',
            )
        names.append(segment)

    if not marked:
        last = names[-1]
        return _Chain(
            [],
            "CATEGORY_MARKER_MISSING",
            f"No (islastsubcategory) marker. Add it to the last name in the "
            f'chain — write "{last}(islastsubcategory)".',
        )
    if len(marked) > 1:
        first, second = names[marked[0]], names[marked[1]]
        return _Chain(
            [],
            "CATEGORY_MARKER_REPEATED",
            f'(islastsubcategory) is on "{first}" and on "{second}". Only the '
            "last name in the chain carries it.",
        )
    if marked[0] != len(names) - 1:
        here, after = names[marked[0]], names[marked[0] + 1]
        return _Chain(
            [],
            "CATEGORY_MARKER_NOT_LAST",
            f'(islastsubcategory) is on "{here}", but "{after}" comes after it. '
            "The marker belongs on the last name in the chain.",
        )

    if len(names) == 1:
        return _Chain(
            [],
            "CATEGORY_LEAF_AT_ROOT",
            f'"{names[0]}" is the only name in the chain, so it would be a '
            "top-level category — and those hold sub-categories, not "
            "products. Add the sub-category this product sits in.",
        )

    return _Chain(names)


def chain_path(segments: list[str]) -> str:
    """`Electronics › Television › Android TV` — the separator every console uses."""
    return " › ".join(segments)


# ── one row ──────────────────────────────────────────────────────────────────


@dataclasses.dataclass
class _Row:
    """One data row, at whatever stage it has reached."""

    number: int
    cells: dict[str, str]
    chain: _Chain | None = None
    node_id: uuid.UUID | None = None
    brand: "VendorBrand | None" = None
    name: str = ""
    service_types: list[str] = dataclasses.field(default_factory=list)
    capacity: str | None = None
    warranty_months: int | None = None
    parameters: list[dict[str, str]] = dataclasses.field(default_factory=list)
    technician_payout_paise: int | None = None
    vendor_price_paise: int | None = None
    #: The live row this one matches, if any — set once the catalogue resolves.
    existing: "ProductModel | None" = None
    #: "create" or "update", set alongside `existing`.
    action: str | None = None
    reject: ProductImportReject | None = None

    def refuse(self, code: str, reason: str, field: str | None = None) -> None:
        if self.reject is not None:
            return  # First reason wins; one row shows one reason.
        raw = self.cells.get(field or "", "") if field else ""
        self.reject = ProductImportReject(
            row=self.number,
            code=code,
            reason=reason,
            field=_BY_KEY[field].header if field and field in _BY_KEY else None,
            value=(raw[:64] or None),
        )


def _parse_service_types(raw: str) -> tuple[list[str], str | None]:
    """Comma-separated subset of `SERVICE_TYPES`. Blank defaults to the baseline."""
    cell = (raw or "").strip()
    if not cell:
        return list(DEFAULT_SERVICE_TYPES), None
    picked = {p.strip() for p in cell.split(",") if p.strip()}
    unknown = sorted(picked - set(SERVICE_TYPES))
    if unknown:
        return [], (
            f"Unknown service type: {', '.join(unknown)}. Choose from "
            f"{', '.join(SERVICE_TYPES)}."
        )
    return [s for s in SERVICE_TYPES if s in picked], None


def _parse_warranty(raw: str) -> tuple[int | None, str | None]:
    cell = (raw or "").strip()
    if not cell:
        return None, None
    try:
        value = int(float(cell))
    except ValueError:
        return None, f'"{cell}" is not a whole number of months.'
    if value < 0 or value > 240:
        return None, "Warranty must be between 0 and 240 months."
    return value, None


def _parse_price(raw: str) -> tuple[int | None, str | None]:
    """A positive paise amount, or `None` if the cell is blank."""
    cell = (raw or "").strip()
    if not cell:
        return None, None
    try:
        value = int(float(cell))
    except ValueError:
        return None, f'"{cell}" is not a whole number of paise.'
    if value <= 0:
        return None, "A price must be a positive number of paise."
    if value > 100_000_000:
        return None, "That is over ₹10,00,000 — check it is paise, not rupees."
    return value, None


_PARAM_SPLIT_RE = re.compile(r"\s*;\s*")


def _parse_parameters(raw: str) -> tuple[list[dict[str, str]], str | None]:
    """`RAM:8 GB; Storage:128 GB` — the same shape the console's spec editor saves."""
    cell = (raw or "").strip()
    if not cell:
        return [], None
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for chunk in _PARAM_SPLIT_RE.split(cell):
        chunk = chunk.strip()
        if not chunk:
            continue
        if ":" not in chunk:
            return [], (
                f'"{chunk}" is not Name:Value. Separate each field\'s name and '
                "value with a colon."
            )
        name, value = chunk.split(":", 1)
        name = name.strip()
        value = value.strip()
        if not name:
            return [], f'"{chunk}" has no field name.'
        if len(name) > 64:
            return [], f'"{name}" is longer than 64 characters.'
        if not value:
            return [], f"Give {name} a value, or remove the field."
        if len(value) > 255:
            return [], f'The value for "{name}" is longer than 255 characters.'
        key = name.casefold()
        if key in seen:
            return [], f"{name} is listed twice — each field needs its own name."
        seen.add(key)
        out.append({"name": name, "value": value})
    if len(out) > MAX_PARAMETERS:
        return [], f"Up to {MAX_PARAMETERS} fields."
    return out, None


def _shape_check(row: _Row, *, is_staff: bool) -> None:
    """Every rule decidable from the row alone — no database, no catalogue."""
    cells = row.cells

    name = cells.get("name", "").strip()
    if not name:
        row.refuse("PRODUCT_NAME_MISSING", "Name the product.", "name")
        return
    if len(name) > 120:
        row.refuse(
            "PRODUCT_NAME_TOO_LONG",
            "The product name is longer than 120 characters.",
            "name",
        )
        return
    row.name = name

    capacity = cells.get("capacity", "").strip()
    if len(capacity) > 64:
        row.refuse(
            "CAPACITY_TOO_LONG", "Capacity is longer than 64 characters.", "capacity"
        )
        return
    row.capacity = capacity or None

    types, err = _parse_service_types(cells.get("serviceTypes", ""))
    if err:
        row.refuse("SERVICE_TYPE_UNKNOWN", err, "serviceTypes")
        return
    row.service_types = types

    warranty, err = _parse_warranty(cells.get("warrantyMonths", ""))
    if err:
        row.refuse("WARRANTY_INVALID", err, "warrantyMonths")
        return
    row.warranty_months = warranty

    params, err = _parse_parameters(cells.get("parameters", ""))
    if err:
        row.refuse("PARAMETER_INVALID", err, "parameters")
        return
    row.parameters = params

    # Prices are not even READ for a vendor's own upload — a vendor has no
    # field to send either figure in anywhere else in this codebase, and a
    # bulk import must not become the one place that changes.
    if is_staff:
        payout, err = _parse_price(cells.get("technicianPayoutPaise", ""))
        if err:
            row.refuse("PRICE_INVALID", err, "technicianPayoutPaise")
            return
        row.technician_payout_paise = payout

        price, err = _parse_price(cells.get("vendorPricePaise", ""))
        if err:
            row.refuse("PRICE_INVALID", err, "vendorPricePaise")
            return
        row.vendor_price_paise = price

        if (row.technician_payout_paise is None) != (row.vendor_price_paise is None):
            row.refuse(
                "PRICE_INCOMPLETE",
                "Give both prices, or leave both blank to keep the current one.",
                "technicianPayoutPaise"
                if row.technician_payout_paise is None
                else "vendorPricePaise",
            )
            return


# ── rows disagreeing with each other ─────────────────────────────────────────


def _reconcile_chains(rows: list[_Row]) -> None:
    """The conflict no single row can see: one chain's leaf is another's parent.

    The same check the ticket importer runs on the same grammar — see its
    docstring for the worked example. Both rows are rejected and the chain is
    dropped from the plan: there is no way to know which was meant.
    """
    leaf_at: dict[tuple[str, ...], _Row] = {}
    for row in rows:
        if row.reject is not None or row.chain is None or not row.chain.segments:
            continue
        leaf_at.setdefault(row.chain.key, row)

    conflicted: dict[tuple[str, ...], tuple[_Row, _Row]] = {}
    for row in rows:
        if row.reject is not None or row.chain is None or not row.chain.segments:
            continue
        key = row.chain.key
        for cut in range(1, len(key)):
            prefix = key[:cut]
            other = leaf_at.get(prefix)
            if other is not None and other is not row:
                conflicted.setdefault(prefix, (other, row))

    for prefix, (leaf_row, deep_row) in conflicted.items():
        assert leaf_row.chain and deep_row.chain
        leaf_name = leaf_row.chain.segments[len(prefix) - 1]
        child_name = deep_row.chain.segments[len(prefix)]
        sentence = (
            f'Row {leaf_row.number} files a product directly in '
            f'"{chain_path(leaf_row.chain.segments)}", and row {deep_row.number} '
            f'puts "{child_name}" under it. "{leaf_name}" can hold products or '
            "sub-categories, not both — decide which and upload again."
        )
        for row in rows:
            if row.reject is not None or row.chain is None:
                continue
            key = row.chain.key
            if key == prefix or (len(key) > len(prefix) and key[: len(prefix)] == prefix):
                row.refuse("CATEGORY_LEAF_CONFLICT", sentence, "category")


_CHAIN_CONFLICTS: dict[str, str] = {
    "NAME_EMPTY": 'There is an empty name in "{path}". Remove the extra separator.',
    "NAME_TOO_LONG": '"{name}" is {length} characters. A category name holds 64 — '
    "shorten it and upload again.",
    "TOO_MANY_SEGMENTS": '"{path}" is {depth} levels deep. The catalogue holds 6, '
    "so shorten the chain and upload again.",
    "TOO_DEEP": '"{name}" would be a seventh level. The catalogue holds 6.',
    "LEAF_AT_ROOT": '"{name}" is the only name in the chain, so it would be a '
    "top-level category — and those hold sub-categories, not products. Add the "
    "sub-category this product sits in.",
    "PARENT_IS_LEAF": '"{name}" is already marked as the last sub-category, so it '
    'holds products and cannot take "{child}" underneath. File this product under '
    '"{name}", or ask for the catalogue to be changed.',
    "HAS_CHILDREN": '"{name}" already has sub-categories under it, so a product '
    "cannot sit directly in it. Name the one this product belongs to.",
    "INACTIVE": '"{name}" is paused, so nothing under it can be ticketed. Ask for '
    "it to be switched back on, or file this under a different category.",
    "RACE_LOST": "Somebody else was changing this category at the same moment. "
    "Upload the file again.",
}


# ── the catalogue, loaded once for the whole file ────────────────────────────

_RESOLVE_CHAINS = text(
    "SELECT chain_ix, segment_ix, resolved_id, segment_name, node_depth, "
    "       action, conflict "
    "  FROM resolve_category_chains("
    "       :company, CAST(:chains AS jsonb), :actor, :dry) "
    " ORDER BY chain_ix, segment_ix"
)


@dataclasses.dataclass
class _ChainResult:
    node_id: uuid.UUID | None = None
    conflict: str | None = None
    at: str = ""
    child: str = ""
    new_segments: list[str] = dataclasses.field(default_factory=list)
    marked_leaf: bool = False
    depth: int = 0


async def _resolve_chains(
    db: AsyncSession,
    company_id: uuid.UUID,
    actor: uuid.UUID | None,
    chains: list[list[str]],
    *,
    dry_run: bool,
) -> list[_ChainResult]:
    """One round trip for every distinct chain in the file.

    The same plpgsql function the ticket importer calls
    (`resolve_category_chains`, migration `c9a41f7b0e83`) — it is a database
    routine shared by any caller that needs this grammar, not a Python
    service the tickets slice owns, so calling it here is not the cross-slice
    import hard rule 4 refuses.
    """
    if not chains:
        return []

    rows = (
        await db.execute(
            _RESOLVE_CHAINS,
            {
                "company": company_id,
                "chains": json.dumps(chains),
                "actor": actor,
                "dry": dry_run,
            },
        )
    ).mappings().all()

    results = [_ChainResult() for _ in chains]
    for record in rows:
        result = results[record["chain_ix"]]
        conflict = record["conflict"]
        if conflict is not None:
            result.conflict = conflict
            result.at = record["segment_name"] or ""
            index = record["segment_ix"]
            source = chains[record["chain_ix"]]
            result.child = source[index + 1] if index + 1 < len(source) else ""
            result.depth = len(source)
            continue
        action = record["action"]
        if action in ("created", "would_create"):
            result.new_segments.append(record["segment_name"] or "")
        elif action in ("marked_leaf", "would_mark_leaf"):
            result.marked_leaf = True
        result.node_id = record["resolved_id"] or result.node_id
        result.depth = len(chains[record["chain_ix"]])
    return results


async def _load_vendor(
    db: AsyncSession, company_id: uuid.UUID, vendor_id: uuid.UUID | None
) -> Vendor:
    """The vendor this sheet's products belong to — through a company-scoped loader.

    404 rather than 403 for anything outside the caller's company, or for a
    `None` id (a vendor account with no vendor membership), like every other
    load-by-id here: a 403 confirms the row exists.
    """
    row = (
        await db.scalar(
            select(Vendor).where(
                Vendor.id == vendor_id,
                Vendor.company_id == company_id,
                Vendor.deleted_at.is_(None),
            )
        )
        if vendor_id is not None
        else None
    )
    if row is None:
        raise HTTPException(status_code=404, detail="Vendor not found")
    if not row.is_active:
        raise _bad_request(
            f"{row.name} is paused, so no products can be added for them."
        )
    return row


# ── the importer ─────────────────────────────────────────────────────────────


async def import_products(
    db: AsyncSession,
    principal: Principal,
    data: bytes,
    filename: str,
    *,
    vendor_id: uuid.UUID | None,
    is_staff: bool,
    dry_run: bool,
    create_categories: bool = False,
    expected_digest: str | None = None,
) -> ProductImportReport:
    """Read the file, and either report what it would do or do it.

    `create_categories` defaults to FALSE, mirroring the ticket importer's own
    confirmation gate — a popup is presentation, and the affirmative has to
    reach the server as a positive assertion the client made.
    """
    digest = hashlib.sha256(data).hexdigest()
    if expected_digest is not None and expected_digest != digest:
        raise AppError(
            409,
            "FILE_CHANGED",
            "That is not the file you just checked. Choose it again so the "
            "numbers you confirm describe what is imported.",
        )

    vendor = await _load_vendor(db, principal.company_id, vendor_id)

    # ── 1. read the sheet ────────────────────────────────────────────────
    raw_rows = list(_iter_rows(data, filename))
    if not raw_rows:
        raise _bad_request("The file is empty")

    header: dict[str, int] = {}
    for index, cell in enumerate(raw_rows[0]):
        key = next((c.key for c in COLUMNS if _norm(cell) in c.aliases), None)
        if key is not None and key not in header:
            header[key] = index
    missing = [_BY_KEY[k].header for k in REQUIRED_COLUMNS if k not in header]
    if missing:
        found = ", ".join(_cell_text(c) for c in raw_rows[0] if _cell_text(c)) or "nothing"
        raise HTTPException(
            status_code=422,
            detail=(
                f"The sheet is missing {', '.join(missing)}. Found: {found}. "
                "Start from the template if you are not sure."
            ),
        )

    body_rows = raw_rows[1:]
    if len(body_rows) > MAX_PRODUCT_IMPORT_ROWS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"That file has {len(body_rows):,} rows. The importer takes "
                f"{MAX_PRODUCT_IMPORT_ROWS} at a time — split it and upload "
                "each part."
            ),
        )

    rows: list[_Row] = []
    for offset, raw in enumerate(body_rows):
        cells = {
            key: _cell_text(raw[i]) if i < len(raw) else ""
            for key, i in header.items()
        }
        if not any(cells.values()):
            continue  # A blank row is spacing, not a mistake.
        rows.append(_Row(number=offset + 2, cells=cells))

    # ── 2. everything decidable from the row alone ───────────────────────
    for row in rows:
        _shape_check(row, is_staff=is_staff)
        if row.reject is None:
            row.chain = parse_category_chain(row.cells.get("category", ""))
            if row.chain.code is not None:
                row.refuse(row.chain.code, row.chain.reason or "", "category")

    # ── 3. rows disagreeing with each other ──────────────────────────────
    _reconcile_chains(rows)

    # ── 4. the catalogue chains, in one call ─────────────────────────────
    live = [r for r in rows if r.reject is None and r.chain is not None]
    distinct: dict[tuple[str, ...], int] = {}
    chain_list: list[list[str]] = []
    for row in live:
        assert row.chain is not None
        if row.chain.key not in distinct:
            distinct[row.chain.key] = len(chain_list)
            chain_list.append(row.chain.segments)

    resolve_dry = dry_run or not create_categories
    chain_results = await _resolve_chains(
        db, principal.company_id, principal.user_id, chain_list, dry_run=resolve_dry
    )

    for row in live:
        assert row.chain is not None
        result = chain_results[distinct[row.chain.key]]
        if result.conflict is not None:
            template = _CHAIN_CONFLICTS.get(
                result.conflict, "That category could not be used."
            )
            row.refuse(
                f"CATEGORY_{result.conflict}",
                template.format(
                    name=result.at,
                    child=result.child,
                    path=chain_path(row.chain.segments),
                    depth=result.depth,
                    length=len(result.at),
                ),
                "category",
            )
        elif result.node_id is None or (result.new_segments and not create_categories):
            row.refuse(
                "CATEGORY_NOT_FOUND",
                f'"{chain_path(row.chain.segments)}" is not in the catalogue '
                "yet. Confirm that it should be created, then import again.",
                "category",
            )
        else:
            row.node_id = result.node_id

    # ── 5. brand, then create-vs-update ──────────────────────────────────
    await _resolve_products(db, principal, vendor, rows, is_staff=is_staff)

    # ── 6. two rows naming the same product ──────────────────────────────
    duplicates = _mark_duplicates(rows)

    # ── 7. what would happen ─────────────────────────────────────────────
    ready = [r for r in rows if r.reject is None]
    to_create = [r for r in ready if r.action == "create"]
    to_update = [r for r in ready if r.action == "update"]

    report = ProductImportReport(
        dryRun=dry_run,
        fileDigest=digest,
        rowsRead=len(rows),
        toCreate=len(to_create),
        toUpdate=len(to_update),
        rejected=sum(1 for r in rows if r.reject is not None),
        rejects=[r.reject for r in rows if r.reject is not None][
            :MAX_PRODUCT_REJECTS_RETURNED
        ],
        categoriesToCreate=_categories_to_create(live, distinct, chain_results),
        duplicateRowsInFile=duplicates,
    )

    if dry_run:
        # Nothing was written, but the ORM may hold objects a later read in
        # this session would otherwise see. Same guard `import_geography` and
        # `tickets.import_service` end a dry run on.
        db.expunge_all()
        return report

    await _commit(db, principal, vendor, to_create, to_update, report, is_staff=is_staff)
    return report


async def _resolve_products(
    db: AsyncSession,
    principal: Principal,
    vendor: Vendor,
    rows: list[_Row],
    *,
    is_staff: bool,
) -> None:
    """Brand, then create-vs-update — for the whole file, in a handful of queries."""
    live = [r for r in rows if r.reject is None and r.node_id is not None]

    # ── brands ───────────────────────────────────────────────────────────
    brands = list(
        await db.scalars(
            select(VendorBrand).where(
                VendorBrand.company_id == principal.company_id,
                VendorBrand.vendor_id == vendor.id,
                VendorBrand.deleted_at.is_(None),
            )
        )
    )
    by_name = {b.name.strip().lower(): b for b in brands}
    approved = [b for b in brands if b.approval_status == APPROVED]
    #: Omitting the brand is allowed only while there is exactly one approved
    #: — the same fallback `masters.service._resolve_brand` makes.
    sole = approved[0] if len(approved) == 1 else None

    for row in live:
        wanted = row.cells.get("brand", "").strip()
        if not wanted:
            if sole is not None:
                row.brand = sole
            elif not approved:
                row.refuse(
                    "BRAND_NONE",
                    "This vendor has no approved brand yet. Add one before "
                    "importing products.",
                    "brand",
                )
            else:
                row.refuse(
                    "BRAND_AMBIGUOUS",
                    "Name the brand. This vendor has more than one approved "
                    "brand, so it cannot be guessed.",
                    "brand",
                )
            continue
        found = by_name.get(wanted.lower())
        if found is None:
            row.refuse(
                "BRAND_NOT_FOUND",
                f'"{wanted}" is not one of {vendor.name}\'s approved brands. '
                "Add it first, or use a brand already on file.",
                "brand",
            )
        elif found.approval_status == PENDING:
            row.refuse(
                "BRAND_NOT_APPROVED",
                f"{found.name} is still waiting for approval, so no product "
                "can carry it yet.",
                "brand",
            )
        elif found.approval_status == REJECTED:
            row.refuse(
                "BRAND_REJECTED",
                f"{found.name} was not approved, so no product can carry it.",
                "brand",
            )
        else:
            row.brand = found

    # ── create vs update, matched on (node, brand, name) ─────────────────
    live = [r for r in rows if r.reject is None and r.brand is not None]
    node_ids = {r.node_id for r in live if r.node_id is not None}
    existing: dict[tuple[uuid.UUID, uuid.UUID, str], ProductModel] = {}
    if node_ids:
        for model in await db.scalars(
            select(ProductModel).where(
                ProductModel.company_id == principal.company_id,
                ProductModel.vendor_id == vendor.id,
                ProductModel.node_id.in_(node_ids),
                ProductModel.deleted_at.is_(None),
            )
        ):
            existing[(model.node_id, model.brand_id, model.name.strip().lower())] = model

    for row in live:
        assert row.node_id is not None and row.brand is not None
        key = (row.node_id, row.brand.id, row.name.lower())
        match = existing.get(key)
        if match is None:
            row.action = "create"
            if is_staff and (
                row.technician_payout_paise is None or row.vendor_price_paise is None
            ):
                row.refuse(
                    "PRICE_MISSING",
                    f'"{row.name}" is new, so it needs both prices to be '
                    "approved — fill in Technician Payout (paise) and Vendor "
                    "Price (paise).",
                    "technicianPayoutPaise"
                    if row.technician_payout_paise is None
                    else "vendorPricePaise",
                )
            continue
        row.existing = match
        row.action = "update"
        if is_staff and match.approval_status != APPROVED:
            waiting = (
                "waiting for approval"
                if match.approval_status == PENDING
                else "was not approved"
            )
            row.refuse(
                "PRODUCT_NOT_APPROVED",
                f"{match.name} is {waiting}. Decide it on Approvals before "
                "updating it from a sheet.",
                "name",
            )


def _mark_duplicates(rows: list[_Row]) -> int:
    """Two rows naming the same (chain, brand, name) — the second is rejected.

    The sheet's OWN repeat, as opposed to a match against a row already in
    the database. Caught here rather than left to the unique index, because a
    lost race there is a 500 where a sentence belongs.
    """
    seen: set[tuple[uuid.UUID, uuid.UUID, str]] = set()
    duplicates = 0
    for row in rows:
        if row.reject is not None or row.node_id is None or row.brand is None:
            continue
        key = (row.node_id, row.brand.id, row.name.lower())
        if key in seen:
            row.refuse(
                "DUPLICATE_IN_FILE",
                f'"{row.name}" is named by an earlier row in this file too. '
                "Each product needs one row.",
                "name",
            )
            duplicates += 1
        else:
            seen.add(key)
    return duplicates


def _categories_to_create(
    live: list[_Row],
    distinct: dict[tuple[str, ...], int],
    results: list[_ChainResult],
) -> list[CategoryToCreateOut]:
    """The confirmation list: every chain with something new in it."""
    counts: dict[tuple[str, ...], int] = {}
    for row in live:
        assert row.chain is not None
        counts[row.chain.key] = counts.get(row.chain.key, 0) + 1

    out: list[CategoryToCreateOut] = []
    for row in live:
        assert row.chain is not None
        key = row.chain.key
        if any(c.path == chain_path(row.chain.segments) for c in out):
            continue
        result = results[distinct[key]]
        if result.conflict is not None:
            continue
        if not result.new_segments and not result.marked_leaf:
            continue
        main = (
            row.chain.segments[CERTIFY_DEPTH]
            if len(row.chain.segments) > CERTIFY_DEPTH
            else None
        )
        out.append(
            CategoryToCreateOut(
                path=chain_path(row.chain.segments),
                newSegments=list(result.new_segments),
                depth=len(row.chain.segments),
                isLeaf=True,
                markedLeaf=result.marked_leaf,
                newMainSubcategory=main is not None and main in result.new_segments,
                rowCount=counts.get(key, 0),
            )
        )
    return out


async def _commit(
    db: AsyncSession,
    principal: Principal,
    vendor: Vendor,
    to_create: list[_Row],
    to_update: list[_Row],
    report: ProductImportReport,
    *,
    is_staff: bool,
) -> None:
    """Write the catalogue and the products, in ONE transaction.

    Categories and products commit together, the same reasoning
    `tickets.import_service._commit` gives: a file that created six
    categories and then failed to write a product would leave a catalogue
    nobody asked for.
    """
    now = datetime.datetime.now(datetime.timezone.utc)
    report.categoriesCreated = sum(
        len(c.newSegments) for c in report.categoriesToCreate
    )

    # ── sort order, one grouped query rather than one per row ────────────
    next_sort: dict[uuid.UUID, int] = {}
    if to_create:
        node_ids = {r.node_id for r in to_create if r.node_id is not None}
        rows = await db.execute(
            select(ProductModel.node_id, func.max(ProductModel.sort_order))
            .where(
                ProductModel.node_id.in_(node_ids),
                ProductModel.deleted_at.is_(None),
            )
            .group_by(ProductModel.node_id)
        )
        next_sort = {node_id: int(m or 0) for node_id, m in rows}

    #: Vendor rows that end up PENDING in this commit — a brand-new product,
    #: or a rejected one the sheet is re-asserting. One bell for the whole
    #: file rather than one per row, the same damping the ticket importer's
    #: pool-frame cap applies for the same reason: a hundred-row sheet must
    #: not ring a hundred times.
    newly_pending = 0

    for row in to_create:
        assert row.node_id is not None and row.brand is not None
        next_sort[row.node_id] = next_sort.get(row.node_id, 0) + 1
        db.add(
            ProductModel(
                company_id=principal.company_id,
                node_id=row.node_id,
                vendor_id=vendor.id,
                brand_id=row.brand.id,
                name=row.name,
                service_types=list(row.service_types),
                capacity=row.capacity,
                warranty_months=row.warranty_months,
                parameters=list(row.parameters),
                image_urls=[],
                technician_payout_paise=(
                    row.technician_payout_paise if is_staff else None
                ),
                vendor_price_paise=(row.vendor_price_paise if is_staff else None),
                approval_status=APPROVED if is_staff else PENDING,
                submitted_at=None if is_staff else now,
                is_active=True,
                sort_order=next_sort[row.node_id],
                created_by=principal.user_id,
            )
        )
        if not is_staff:
            newly_pending += 1

    for row in to_update:
        model = row.existing
        assert model is not None
        # The name may differ only in case from what is stored — the match key
        # is case-insensitive — so this still picks up a capitalisation fix.
        model.name = row.name
        model.service_types = list(row.service_types)
        model.capacity = row.capacity
        model.warranty_months = row.warranty_months
        model.parameters = list(row.parameters)
        if is_staff:
            # Both-or-neither was already enforced in `_shape_check`: either
            # both prices are set, repricing it, or both are None, leaving the
            # live figures untouched.
            if row.technician_payout_paise is not None:
                model.technician_payout_paise = row.technician_payout_paise
            if row.vendor_price_paise is not None:
                model.vendor_price_paise = row.vendor_price_paise
        elif model.approval_status == REJECTED:
            # The vendor's bulk twin of `update_own_model`'s resubmission: a
            # rejected product the sheet still names is read as the vendor
            # answering the refusal, so it goes back for a decision. Prices
            # are left alone — there are none yet to clear.
            model.approval_status = PENDING
            model.submitted_at = now
            model.decided_at = None
            model.decided_by = None
            model.rejection_reason = None
            newly_pending += 1
        model.updated_by = principal.user_id

    report.created = len(to_create)
    report.updated = len(to_update)

    raised: uuid.UUID | None = None
    if newly_pending:
        raised = await _notify_bulk_submitted(db, principal, vendor, newly_pending)

    await db.commit()

    if raised is not None:
        await publish_notification(
            db,
            company_id=principal.company_id,
            pincode=None,
            notification_id=raised,
        )


async def _notify_bulk_submitted(
    db: AsyncSession, principal: Principal, vendor: Vendor, count: int
) -> uuid.UUID:
    """Tell staff there is something to price — once for the whole file.

    `pincode=None`, the same reasoning `masters.service._notify_submitted`
    gives: a catalogue is company-wide, so there is no place to scope this to,
    and the cost (every staff reader hears it, including an Area Manager who
    cannot open it) is the same one already accepted there.
    """
    raised = await notify(
        db,
        company_id=principal.company_id,
        kind="product_submitted",
        title=f"{vendor.name} submitted {count} product{'' if count == 1 else 's'}",
        detail="From a spreadsheet import. Set both prices to approve them.",
        to="/approvals",
    )
    return raised.id
