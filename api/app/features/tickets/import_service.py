"""The bulk ticket importer: a spreadsheet in, tickets out.

The third member of a family — `geo.import_geography` and
`masters.import_serials` are the other two — and it keeps their contract
exactly:

  * **Two passes over one file.** A dry run writes nothing and returns exactly
    what the commit would do, so every number the console shows is the server's
    own count rather than a guess made in a browser that never parsed the file.
  * **Rejected rows never block the good ones.** A bad cell costs that row, not
    the upload. The count is exact; the list is capped.
  * **Nothing is parsed in the client.** No spreadsheet library ships to a
    browser or an app.

What is new here, and what the rest of this docstring is about, is that a ticket
needs a whole CATALOGUE to exist before it can be raised.

## The Category column names a chain, and the importer builds it

    Electronics, Television, Android TV(islastsubcategory)

Comma-separated names, root first, the last one marked. The marker is the
sheet's spelling of the console's "This is the last sub-category" checkbox —
which is what sets `is_leaf`, and therefore which level holds products.

⚠ **A comma is not a safe delimiter on its own.** `product_nodes.name` is free
text, so `Cooking, Baking & Grills` is a legal category name and no parser can
tell it from two names. So: **if a cell contains `>`, `>` is the delimiter;
otherwise the comma is.** Per cell, never mixed, never guessed. A category whose
name contains `>` cannot be expressed, which is a trade worth making — the
console renders paths with `›` (U+203A), so nothing on screen collides.

The marker is strictly a checksum. The last segment is the leaf by definition,
so it carries no information the parser could not derive. It is required anyway,
because a truncated paste — `Electronics, Television` where `Android TV` was
meant — otherwise silently builds a branch in the wrong shape, and a vendor
cannot undo it: there is no vendor DELETE for a node.

Resolving and creating the chain is `resolve_category_chains`, a plpgsql
function (migration `c9a41f7b0e83`). The whole file goes in one call. It returns
per-segment rows with a `conflict` code rather than raising, which is what lets
one bad chain cost one row.

## A product the file names but the catalogue lacks is SUBMITTED, not ticketed

It is created `pending` with both prices NULL, because a vendor may never price
a product — and `approved_is_priced` means an approved row must carry both
figures, so this holds no matter who runs the import: staff cannot price from a
sheet with no price columns either. The rows naming it are then rejected with
"waiting to be priced". Somebody approves it on `/approvals`, the same file goes
in again, and those rows import.

## Which is why `Reference` exists

That second upload re-presents every row that already worked. Without a handle
they would be raised again and CHARGED again, and nothing already on `tickets`
could catch it: `serial_number` is deliberately not unique, because a service
call on a unit installed months ago repeats it. So a row carries the vendor's
own job number, `tickets.external_ref`, unique per vendor — and a repeat is
SKIPPED and reported, never rejected and never charged.

## What the importer deliberately does not put in the sheet

  * **No `Vendor` column.** On the portal the vendor is the principal; for staff
    it is a form field resolved through a company-scoped loader. A vendor id in
    a file a vendor can edit is the tenancy boundary sitting in a request body,
    which is the reason `TicketCreateRequest` has no such field either.
  * **No `Latitude` / `Longitude`.** `TicketCreateRequest` says it outright:
    null is the true statement "this address was typed, not picked", and it is
    what every Excel and API ticket will always say. These keep the pincode
    proof rule.
  * **No slot columns.** A slot must be one of the windows `offered_slots`
    produces, with a 90-minute lead, and the server refuses one past the service
    level. A sheet authored at 09:00 and uploaded at 14:00 would have half its
    slots refused for a reason nobody could fix by editing the file. Every
    imported ticket lands `Slot Pending` — exactly what the intake form already
    does when the slot box is left empty — and the customer picks.

## Shape rules are not restated here

Every rule that can be decided from a row alone is enforced by building a real
`TicketCreateRequest` and catching `ValidationError`. That is what keeps the
sheet and the form incapable of disagreeing about a description, a service
level, a phone or a date. The two rules that need a row from the database — the
serial and the pincode — reuse the single-ticket path's own sentences through
`service.serial_refusal_text` and `service.pincode_refusal_text`.
"""

import csv
import dataclasses
import datetime
import hashlib
import io
import json
import re
import secrets
import uuid
from collections.abc import Iterator

from fastapi import HTTPException, status
from pydantic import ValidationError
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.credits import charge_tickets, quote_tickets
from app.core.vendor_credits import assert_within_limit as assert_vendor_within_limit
from app.core.vendor_credits import quote as quote_vendor_credit
from app.core.deps import Principal
from app.core.errors import AppError
from app.core.product_tree import APPROVED, PENDING, REJECTED
from app.core.realtime import publish_pool_changed, publish_ticket_changed
from app.core.rules import resolve_rules
from app.core.sequences import allocate
from app.core.service_types import SERVICE_TYPES
from app.core.slots import IST
from app.core.tickets import (
    DEFAULT_SERVICE_LEVEL_HOURS,
    SERVICE_LEVEL_HOURS,
)
from app.features.tickets.schemas import (
    CategoryToCreate,
    ProductToSubmit,
    TicketCreateRequest,
    TicketImportReject,
    TicketImportReport,
)
from app.models.product import ProductModel, ProductModelSerial, ProductNode
from app.models.territory import Pincode
from app.models.ticket import Ticket
from app.models.vendor import Vendor
from app.models.vendor_brand import VendorBrand

#: 500 tickets is the ceiling, and it is about the WRITE rather than the parse.
#: Each row stamps prices and a rules snapshot, charges credits, writes an event
#: and queues a WhatsApp; the whole commit is one transaction holding the
#: company's credit lock. Bigger files are refused with a sentence that says to
#: split them, which is honest, rather than accepted and timed out.
MAX_TICKET_IMPORT_ROWS = 500

#: Generous for 500 rows of text, and the same read-with-a-ceiling guard both
#: siblings use — the declared size is not trusted.
MAX_TICKET_UPLOAD_BYTES = 8 * 1024 * 1024

#: How many rejects travel back. The COUNT is always exact; this caps the list
#: so a catastrophically wrong file cannot return a 40 MB response.
MAX_TICKET_REJECTS_RETURNED = 200

#: Above this many tickets, per-ticket realtime frames are not emitted: 500
#: `pg_notify` calls would swamp every open console, and the importer's own
#: response is what the uploader is looking at anyway. Pool frames are deduped
#: rather than dropped, because a technician's app has nothing else to tell it.
MAX_PER_TICKET_FRAMES = 25

#: The catalogue level a technician certifies on. Imported here rather than
#: from `core.product_tree` only to keep the reason next to the warning it
#: drives: a brand-new node at this depth has NOBODY certified on it, so its
#: tickets escalate the moment they are raised.
CERTIFY_DEPTH = 1

_SHEET_NAME = "Tickets"
_README_SHEET = "How to fill this in"

#: The zero UUID is never a real node; used as a stand-in while a row's shape is
#: validated, before the catalogue has resolved anything.
_PLACEHOLDER_ID = uuid.UUID("00000000-0000-0000-0000-000000000000")

_MARKER_RE = re.compile(r"\(([^()]*)\)\s*$")


# ── the template ─────────────────────────────────────────────────────────────


@dataclasses.dataclass(frozen=True)
class _Column:
    key: str
    header: str
    #: Everything that normalises to one of these maps to this column, so
    #: `Pin Code`, `pincode` and `PINCODE` all land.
    aliases: tuple[str, ...]
    width: int
    example_a: str = ""
    example_b: str = ""


def _norm(raw: object) -> str:
    """Lowercased, letters and digits only — the header matcher both siblings use."""
    return re.sub(r"[^a-z0-9]", "", str(raw or "").lower())


COLUMNS: tuple[_Column, ...] = (
    _Column(
        "serialNumber",
        "Serial Number",
        ("serialnumber", "serial", "serialno", "serialnos", "sno", "srno"),
        22,
        "SN-EXAMPLE-000001",
        "SN-EXAMPLE-000002",
    ),
    _Column(
        "category",
        "Category",
        ("category", "categories", "categorychain", "categorypath"),
        52,
        "Electronics, Television, Android TV(islastsubcategory)",
        "Home Appliances, Washing Machine(islastsubcategory)",
    ),
    _Column("brand", "Brand", ("brand", "brandname"), 16, "Meridian", "Sunview"),
    _Column(
        "model",
        "Model",
        ("model", "modelname", "product", "productmodel", "productname"),
        24,
        '43 inch 4K UHD',
        "7 kg Front Load",
    ),
    _Column(
        "serviceType",
        "Service Type",
        ("servicetype", "type", "requesttype"),
        20,
        "Installation + Demo",
        "Service",
    ),
    _Column(
        "description",
        "Problem Description",
        ("problemdescription", "description", "problem", "issue"),
        40,
        "",
        "Drum rattles on spin and water drains slowly",
    ),
    _Column(
        "customerName",
        "Customer Name",
        ("customername", "customer", "name"),
        20,
        "Anita Rao",
        "Imran Sheikh",
    ),
    _Column(
        "customerPhone",
        "Mobile Number",
        ("mobilenumber", "mobile", "phone", "phonenumber", "contact", "contactnumber"),
        18,
        "+919876543210",
        "+919812345678",
    ),
    _Column(
        "address",
        "Address",
        ("address", "customeraddress", "addressline"),
        34,
        "Flat 402, Palm Grove, Baner",
        "12 MG Road",
    ),
    _Column("city", "City", ("city", "town"), 16, "Pune", "Hyderabad"),
    _Column("state", "State", ("state",), 16, "Maharashtra", "Telangana"),
    _Column(
        "pincode",
        "Pin Code",
        ("pincode", "pin", "pincodes", "postalcode", "zip"),
        12,
        "411045",
        "500001",
    ),
    _Column(
        "expectedDate",
        "Expected Date",
        ("expecteddate", "expected", "date", "preferreddate"),
        16,
        "28-09-2026",
        "30-09-2026",
    ),
    _Column(
        "serviceLevelHours",
        "Service Level (hours)",
        ("servicelevelhours", "servicelevel", "sla", "slahours"),
        18,
        "24",
        "48",
    ),
    _Column(
        "externalRef",
        "Reference",
        ("reference", "ref", "ordernumber", "orderno", "jobnumber", "docket"),
        18,
        "PO-2026-0001",
        "PO-2026-0002",
    ),
)

#: The columns a file must carry. `Problem Description`, `Service Level (hours)`
#: and `Reference` are the three that may be absent entirely.
REQUIRED_COLUMNS = tuple(
    c.key
    for c in COLUMNS
    if c.key not in ("description", "serviceLevelHours", "externalRef")
)

_BY_KEY = {c.key: c for c in COLUMNS}


def build_ticket_template() -> io.BytesIO:
    """The starter file: one sheet of columns, one sheet of rules.

    The two example rows are deliberately different depths — a three-level chain
    and a two-level one — because "chains may be different lengths" is the thing
    about this format people get wrong, and showing it beats saying it.
    """
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = _SHEET_NAME
    sheet.append([c.header for c in COLUMNS])
    sheet.append([c.example_a for c in COLUMNS])
    sheet.append([c.example_b for c in COLUMNS])
    for i, col in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[
            openpyxl.utils.get_column_letter(i)
        ].width = col.width

    notes = book.create_sheet(_README_SHEET)
    for line in _README_LINES:
        notes.append([line])
    notes.column_dimensions["A"].width = 110

    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


_README_LINES: tuple[str, ...] = (
    "How to fill this in",
    "",
    "Every row is one ticket. Fill in the Tickets sheet and upload it.",
    "Rows that cannot be read are listed with a reason and do not stop the rest.",
    "",
    "Category",
    "  Name the whole chain, from the top-level category down to the level the",
    "  product is filed under, separated by commas:",
    "",
    "      Electronics, Television, Android TV(islastsubcategory)",
    "",
    "  Mark the LAST name with (islastsubcategory). That is the level products",
    "  sit in. Anything in the chain that does not exist yet is created — you are",
    "  shown exactly what will be created before anything is saved.",
    "  A chain can be up to 6 names long, and each name up to 64 characters.",
    "  Different rows may name chains of different lengths.",
    "",
    "  If a category name contains a comma of its own, separate the names with >",
    "  instead. Then every comma in the cell is part of a name:",
    "",
    "      Home, Cooking, Baking > Ovens(islastsubcategory)     <- wrong",
    "      Home > Cooking, Baking > Ovens(islastsubcategory)    <- right",
    "",
    "Brand and Model",
    "  Brand must be one of your approved brands. Leave it blank only if you have",
    "  exactly one. A Model that is not in the catalogue is added for you and sent",
    "  for approval — it has no price yet, so that row does not import today.",
    "  Upload the file again once it has been approved.",
    "",
    "Problem Description",
    "  Leave it blank for Installation + Demo. Required, and at least 10",
    "  characters, for Tech Visit and Service.",
    "",
    "Expected Date",
    "  Write it as 28-09-2026 or 2026-09-28. It cannot be in the past.",
    "",
    "Service Level (hours)",
    "  12, 24, 36 or 48. Blank means 24.",
    "",
    "Reference",
    "  Your own order or job number. It is what makes it safe to upload the same",
    "  file twice: a row whose reference is already on a ticket is skipped, not",
    "  raised again. Leave it blank and that protection does not apply.",
    "",
    "The customer picks the time",
    "  There is no slot column. Every imported ticket is raised as Slot Pending",
    "  and the customer is sent a link to choose a window.",
)


# ── reading the file ─────────────────────────────────────────────────────────


def _bad_request(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=detail)


def _iter_rows(data: bytes, filename: str) -> Iterator[tuple]:
    """Yield raw tuples from .xlsx or .csv, streaming in both cases.

    A near-twin of `geo.service._iter_rows` and `masters.service._iter_serial_rows`,
    kept here for the reason theirs gives: hard rule 4 forbids one slice
    importing another's service, and sharing fifteen lines is not worth coupling
    tickets to the product master.

    ⚠ One divergence, and it matters. Both siblings read
    `book[book.sheetnames[0]]` because their templates have one sheet. This
    template has two, so it PREFERS the sheet named `Tickets` — otherwise
    somebody reordering the tabs would have the instructions parsed as data.
    """
    if filename.lower().endswith(".csv"):
        decoded = data.decode("utf-8-sig", errors="replace")
        yield from csv.reader(io.StringIO(decoded))
        return

    try:
        import openpyxl
    except ModuleNotFoundError:  # pragma: no cover - dependency is in requirements
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Spreadsheet support is not installed on this server",
        )
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise _bad_request("That file could not be opened as a spreadsheet")
    try:
        name = next(
            (s for s in book.sheetnames if _norm(s) == _norm(_SHEET_NAME)),
            book.sheetnames[0],
        )
        yield from book[name].iter_rows(values_only=True)
    finally:
        book.close()


def _cell_text(value: object) -> str:
    """One spreadsheet cell as the string a human meant by it.

    ⚠ The float branch is the one that matters, and it matters more here than in
    either sibling because three columns can trip it. A serial, a pincode and a
    mobile number are all digits, and openpyxl hands back `500001.0` and
    `9.1987654321e+11` for cells somebody typed as numbers. Both siblings carry
    the identical guard; it is the most valuable line in any of the three.
    """
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


#: `28-09-2026` first: the Indian order is what a vendor here will type, and the
#: template's example says so. ISO second, because a date-formatted cell that
#: `_cell_text` already turned into a string arrives that way.
_DATE_FORMATS = ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d")


def _cell_date(raw: str) -> datetime.date | None:
    """A date, or None if it cannot be read as one.

    ⚠ `03-04-2026` is read as 3 April, and there is no way to know whether
    somebody meant 4 March. No parser can fix that; the template's example and
    the Read-me line are the mitigation, and a real date-formatted column in the
    sheet avoids it entirely because openpyxl hands back a `datetime`.
    """
    value = (raw or "").strip()
    if not value:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


# ── the Category chain ───────────────────────────────────────────────────────


def _is_marker(token: str) -> bool:
    return _norm(token) == "islastsubcategory"


def _looks_like_marker(token: str) -> bool:
    """A parenthesised token that was MEANT to be the marker but is not it.

    Worth telling apart from an ordinary bracketed word — `Television (Large)`
    is a legal name — so a misspelling gets "write it exactly as …" rather than
    being silently accepted as part of the name.
    """
    normalised = _norm(token)
    return not _is_marker(normalised) and "last" in normalised and "sub" in normalised


@dataclasses.dataclass
class _Chain:
    """One Category cell, parsed."""

    segments: list[str]
    #: Set when the cell cannot be read at all. The row is rejected with this.
    code: str | None = None
    reason: str | None = None

    @property
    def key(self) -> tuple[str, ...]:
        return tuple(s.lower() for s in self.segments)


def parse_chain(raw: str) -> _Chain:
    """Split a Category cell into names and check the marker.

    Everything answerable from the cell alone. Depth, leaf conflicts and
    anything needing a row from the database are `resolve_category_chains`'s
    job.
    """
    cell = (raw or "").strip()
    if not cell:
        return _Chain(
            [],
            "CATEGORY_EMPTY",
            "No category. Name the chain, ending with the last sub-category.",
        )

    # `>` wins when present, so a category name may contain a comma. Per cell,
    # never mixed — see the module docstring.
    delimiter = ">" if ">" in cell else ","
    raw_segments = cell.split(delimiter)

    if raw_segments and not raw_segments[0].strip():
        return _Chain(
            [],
            "CATEGORY_LEADING_DELIMITER",
            f'The chain starts with "{delimiter}". Remove it, or add the '
            "category before it.",
        )
    if raw_segments and not raw_segments[-1].strip():
        return _Chain(
            [],
            "CATEGORY_TRAILING_DELIMITER",
            f'The chain ends with "{delimiter}". Remove it, or add the missing '
            "sub-category after it.",
        )

    names: list[str] = []
    marked: list[int] = []
    for index, piece in enumerate(raw_segments):
        segment = re.sub(r"\s+", " ", piece.strip())
        found = _MARKER_RE.search(segment)
        if found is not None:
            inner = found.group(1)
            if _is_marker(inner):
                marked.append(index)
                segment = segment[: found.start()].strip()
            elif _looks_like_marker(inner):
                return _Chain(
                    [],
                    "CATEGORY_MARKER_MISSPELT",
                    f'"({inner})" is not the marker. Write it exactly as '
                    "(islastsubcategory).",
                )
        if not segment:
            if index in marked:
                return _Chain(
                    [],
                    "CATEGORY_MARKER_NO_NAME",
                    "The marker has no name in front of it. Write it on the "
                    'last sub-category, like "Android TV(islastsubcategory)".',
                )
            neighbours = [n for n in (names[-1] if names else None,) if n]
            between = f' after "{neighbours[0]}"' if neighbours else ""
            return _Chain(
                [],
                "CATEGORY_EMPTY_LEVEL",
                f"There is an empty name{between}. Remove the extra "
                f'"{delimiter}".',
            )
        names.append(segment)

    if not marked:
        last = names[-1]
        return _Chain(
            [],
            "CATEGORY_MARKER_MISSING",
            f"No (islastsubcategory) marker. Add it to the last name in the "
            f'chain — write "{last}(islastsubcategory)".',
        )
    if len(marked) > 1:
        first, second = names[marked[0]], names[marked[1]]
        return _Chain(
            [],
            "CATEGORY_MARKER_REPEATED",
            f'(islastsubcategory) is on "{first}" and on "{second}". Only the '
            "last name in the chain carries it.",
        )
    if marked[0] != len(names) - 1:
        here, after = names[marked[0]], names[marked[0] + 1]
        return _Chain(
            [],
            "CATEGORY_MARKER_NOT_LAST",
            f'(islastsubcategory) is on "{here}", but "{after}" comes after it. '
            "The marker belongs on the last name in the chain.",
        )

    # A one-name chain asks for a ROOT that holds products, which the
    # `leaf_below_root` CHECK forbids outright.
    #
    # ⚠ Caught HERE and not left to the database, and that is not merely an
    # earlier message. An unrejected one-name chain is a leaf path of length 1,
    # so `_reconcile_chains` would read every ordinary `Electronics, Television`
    # row in the file as disagreeing with it — one impossible cell would reject
    # the whole sheet with a sentence about a conflict that is not the problem.
    if len(names) == 1:
        return _Chain(
            [],
            "CATEGORY_LEAF_AT_ROOT",
            f'"{names[0]}" is the only name in the chain, so it would be a '
            "top-level category — and those hold sub-categories, not products. "
            "Add the sub-category this product sits in.",
        )

    return _Chain(names)


def chain_path(segments: list[str]) -> str:
    """`Electronics › Television › Android TV` — the separator both consoles use."""
    return " › ".join(segments)


# ── one row ──────────────────────────────────────────────────────────────────


#: Which column a Pydantic field belongs to, so a reject can say where to look.
_FIELD_COLUMN = {
    "serialNumber": "serialNumber",
    "serviceType": "serviceType",
    "description": "description",
    "customerName": "customerName",
    "customerPhone": "customerPhone",
    "address": "address",
    "city": "city",
    "state": "state",
    "pincode": "pincode",
    "expectedDate": "expectedDate",
    "serviceLevelHours": "serviceLevelHours",
}


@dataclasses.dataclass
class _Row:
    """One data row, at whatever stage it has reached."""

    number: int
    cells: dict[str, str]
    chain: _Chain | None = None
    body: TicketCreateRequest | None = None
    reject: TicketImportReject | None = None
    #: Filled once the catalogue has resolved.
    node_id: uuid.UUID | None = None
    model_id: uuid.UUID | None = None
    brand: "VendorBrand | None" = None
    model: "ProductModel | None" = None
    #: Set when this row names a product the catalogue does not hold yet.
    product_plan: "_ProductPlan | None" = None
    #: A row whose Reference is already on a ticket: skipped, never rejected.
    already: bool = False

    def refuse(self, code: str, reason: str, field: str | None = None) -> None:
        if self.reject is not None:
            return  # First reason wins; one row shows one reason.
        raw = self.cells.get(field or "", "") if field else ""
        self.reject = TicketImportReject(
            row=self.number,
            code=code,
            reason=reason,
            field=_BY_KEY[field].header if field and field in _BY_KEY else None,
            value=(raw[:64] or None),
        )


def _pydantic_reason(err: dict) -> tuple[str, str | None]:
    """Turn one Pydantic error into a sentence and the column it belongs to.

    Messages raised by `TicketCreateRequest`'s own `model_validator` are already
    house sentences — "Describe the problem — a Tech Visit needs to say what is
    wrong…" — and are passed through verbatim, minus Pydantic's own prefix. Only
    a FIELD failure gets dressed, and then only enough to name the column.
    """
    location = [str(p) for p in err.get("loc", ())]
    field = next((p for p in location if p in _FIELD_COLUMN), None)
    message = str(err.get("msg", "")).removeprefix("Value error, ").strip()
    if field is None:
        return (message or "This row could not be read.", None)
    header = _BY_KEY[field].header
    return (f"{header}: {message}", field)


def _shape_check(row: _Row, today_ist: datetime.date) -> None:
    """Every rule decidable from the row alone, enforced by the real schema.

    The ids are placeholders — the catalogue has not resolved yet — which is the
    whole trick: substituting them lets `TicketCreateRequest` adjudicate the
    description rule, the service type, the service level, the phone and the
    pincode shape, so none of them is restated in this module and none of them
    can drift from the form.

    ⚠ **The past-date rule is the one exception, and it has to be.** It is NOT
    in the schema — `create_ticket` applies it separately, because "has this day
    gone" needs a clock and a schema has none. Restating it here is therefore
    unavoidable; what is avoidable is restating it DIFFERENTLY, so the sentence
    is the one `create_ticket` raises, word for word.

    `today_ist` is computed once for the whole file and passed in, so every row
    is judged against one instant. Judged in IST, not UTC: for five and a half
    hours every evening the two disagree about what day it is, and the person
    who filled this sheet in is in the former.
    """
    cells = row.cells

    service_level_raw = cells.get("serviceLevelHours", "").strip()
    if service_level_raw:
        try:
            service_level = int(float(service_level_raw))
        except ValueError:
            row.refuse(
                "SERVICE_LEVEL_INVALID",
                f'"{service_level_raw}" is not a service level we offer. Use '
                + ", ".join(str(h) for h in SERVICE_LEVEL_HOURS)
                + ".",
                "serviceLevelHours",
            )
            return
    else:
        service_level = DEFAULT_SERVICE_LEVEL_HOURS

    expected_raw = cells.get("expectedDate", "")
    expected = _cell_date(expected_raw)
    if expected is None:
        row.refuse(
            "DATE_INVALID",
            f'"{expected_raw}" is not a date we can read. Write it as '
            "28-09-2026 or 2026-09-28.",
            "expectedDate",
        )
        return
    if expected < today_ist:
        row.refuse(
            "DATE_IN_PAST",
            "The expected date has already passed — pick today or later.",
            "expectedDate",
        )
        return

    description = cells.get("description", "").strip() or None
    try:
        row.body = TicketCreateRequest(
            subcategoryId=_PLACEHOLDER_ID,
            modelId=_PLACEHOLDER_ID,
            serviceType=cells.get("serviceType", "").strip(),
            description=description,
            serialNumber=cells.get("serialNumber", "").strip(),
            customerName=cells.get("customerName", "").strip(),
            customerPhone=cells.get("customerPhone", "").strip(),
            address=cells.get("address", "").strip(),
            city=cells.get("city", "").strip(),
            state=cells.get("state", "").strip(),
            pincode=cells.get("pincode", "").strip(),
            expectedDate=expected,
            serviceLevelHours=service_level,
        )
    except ValidationError as exc:
        reason, field = _pydantic_reason(exc.errors()[0])
        row.refuse("ROW_INVALID", reason, field)


# ── rows disagreeing with each other ─────────────────────────────────────────


def _reconcile_chains(rows: list[_Row]) -> None:
    """The conflict no single row can see: one chain's leaf is another's parent.

    Row 4 says `Electronics, Television(islastsubcategory)` and row 9 says
    `Electronics, Television, Android TV(islastsubcategory)`. `Television` would
    have to hold products AND sub-categories, and `is_leaf` is exactly the flag
    that says it cannot be both.

    **Both rows are rejected, and the chain is dropped from the plan.** There is
    no way to know which was meant, and creating `Television` as a leaf because
    row 4 came first would build the branch one of them did not ask for — while
    the other's rows fail for a reason that no longer matches what is on screen.

    This is the shape the requirement's own example asks about: rows of
    different depths in one file. Detected entirely in memory, before anything
    is written.
    """
    leaf_at: dict[tuple[str, ...], _Row] = {}
    for row in rows:
        if row.reject is not None or row.chain is None or not row.chain.segments:
            continue
        leaf_at.setdefault(row.chain.key, row)

    conflicted: dict[tuple[str, ...], tuple[_Row, _Row]] = {}
    for row in rows:
        if row.reject is not None or row.chain is None or not row.chain.segments:
            continue
        key = row.chain.key
        # Every proper prefix of this chain is a node that must hold children.
        for cut in range(1, len(key)):
            prefix = key[:cut]
            other = leaf_at.get(prefix)
            if other is not None and other is not row:
                conflicted.setdefault(prefix, (other, row))

    for prefix, (leaf_row, deep_row) in conflicted.items():
        assert leaf_row.chain and deep_row.chain
        leaf_name = leaf_row.chain.segments[len(prefix) - 1]
        child_name = deep_row.chain.segments[len(prefix)]
        sentence = (
            f'Row {leaf_row.number} files products directly in '
            f'"{chain_path(leaf_row.chain.segments)}", and row {deep_row.number} '
            f'puts "{child_name}" under it. "{leaf_name}" can hold products or '
            "sub-categories, not both — decide which and upload again."
        )
        for row in rows:
            if row.reject is not None or row.chain is None:
                continue
            key = row.chain.key
            if key == prefix or (len(key) > len(prefix) and key[: len(prefix)] == prefix):
                row.refuse("CATEGORY_LEAF_CONFLICT", sentence, "category")


#: What `resolve_category_chains` can say went wrong, and what a person reads.
#: `{name}` is the segment it stopped on; `{child}` the one that could not go
#: under it. Kept beside the function's own codes so the two are read together.
_CHAIN_CONFLICTS: dict[str, str] = {
    "NAME_EMPTY": 'There is an empty name in "{path}". Remove the extra separator.',
    "NAME_TOO_LONG": '"{name}" is {length} characters. A category name holds 64 — '
    "shorten it and upload again.",
    "TOO_MANY_SEGMENTS": '"{path}" is {depth} levels deep. The catalogue holds 6, '
    "so shorten the chain and upload again.",
    "TOO_DEEP": '"{name}" would be a seventh level. The catalogue holds 6.',
    "LEAF_AT_ROOT": '"{name}" is the only name in the chain, so it would be a '
    "top-level category — and those hold sub-categories, not products. Add the "
    "sub-category this product sits in.",
    "PARENT_IS_LEAF": '"{name}" is already marked as the last sub-category, so it '
    'holds products and cannot take "{child}" underneath. File this product under '
    '"{name}", or ask for the catalogue to be changed.',
    "HAS_CHILDREN": '"{name}" already has sub-categories under it, so a product '
    "cannot sit directly in it. Name the one this product belongs to.",
    "INACTIVE": '"{name}" is paused, so nothing under it can be ticketed. Ask for '
    "it to be switched back on, or file this under a different category.",
    "RACE_LOST": "Somebody else was changing this category at the same moment. "
    "Upload the file again.",
}


# ── the catalogue, loaded once for the whole file ────────────────────────────


_RESOLVE_CHAINS = text(
    "SELECT chain_ix, segment_ix, resolved_id, segment_name, node_depth, "
    "       action, conflict "
    "  FROM resolve_category_chains("
    "       :company, CAST(:chains AS jsonb), :actor, :dry) "
    " ORDER BY chain_ix, segment_ix"
)


@dataclasses.dataclass
class _ChainResult:
    """What the function said about one distinct chain."""

    node_id: uuid.UUID | None = None
    conflict: str | None = None
    #: The segment the conflict stopped on, and the one below it.
    at: str = ""
    child: str = ""
    #: Names this import would create, in order, root first.
    new_segments: list[str] = dataclasses.field(default_factory=list)
    #: The existing node that will be ticked as the last sub-category.
    marked_leaf: bool = False
    depth: int = 0


async def _resolve_chains(
    db: AsyncSession,
    company_id: uuid.UUID,
    actor: uuid.UUID | None,
    chains: list[list[str]],
    *,
    dry_run: bool,
) -> list[_ChainResult]:
    """One round trip for every distinct chain in the file.

    The whole point of the plpgsql function: a 500-row sheet naming eight chains
    is eight chains' worth of work in one call, instead of up to 3,000
    individual probes.
    """
    if not chains:
        return []

    rows = (
        await db.execute(
            _RESOLVE_CHAINS,
            {
                "company": company_id,
                "chains": json.dumps(chains),
                "actor": actor,
                "dry": dry_run,
            },
        )
    ).mappings().all()

    results = [_ChainResult() for _ in chains]
    for record in rows:
        result = results[record["chain_ix"]]
        conflict = record["conflict"]
        if conflict is not None:
            result.conflict = conflict
            result.at = record["segment_name"] or ""
            index = record["segment_ix"]
            source = chains[record["chain_ix"]]
            result.child = source[index + 1] if index + 1 < len(source) else ""
            result.depth = len(source)
            continue
        action = record["action"]
        if action in ("created", "would_create"):
            result.new_segments.append(record["segment_name"] or "")
        elif action in ("marked_leaf", "would_mark_leaf"):
            result.marked_leaf = True
        # The LAST segment's id is the node products hang off.
        result.node_id = record["resolved_id"] or result.node_id
        result.depth = len(chains[record["chain_ix"]])
    return results


async def _load_vendor(
    db: AsyncSession, company_id: uuid.UUID, vendor_id: uuid.UUID
) -> Vendor:
    """The vendor these tickets belong to — through a company-scoped loader.

    404 rather than 403 for anything outside the caller's company, like every
    other load-by-id here: a 403 confirms the row exists.
    """
    row = await db.scalar(
        select(Vendor).where(
            Vendor.id == vendor_id,
            Vendor.company_id == company_id,
            Vendor.deleted_at.is_(None),
        )
    )
    if row is None:
        raise HTTPException(status_code=404, detail="Vendor not found")
    if not row.is_active:
        raise _bad_request(
            f"{row.name} is paused, so no tickets can be raised for them."
        )
    return row


# ── the importer ─────────────────────────────────────────────────────────────


async def import_tickets(
    db: AsyncSession,
    principal: Principal,
    data: bytes,
    filename: str,
    *,
    vendor_id: uuid.UUID,
    dry_run: bool,
    create_categories: bool = False,
    submit_products: bool = False,
    expected_digest: str | None = None,
    show_credits: bool = False,
) -> TicketImportReport:
    """Read the file, and either report what it would do or do it.

    `create_categories` and `submit_products` default to FALSE, and that is the
    server's half of the confirmation. A popup is presentation; the affirmative
    has to reach the server as a positive assertion the client made, or the
    guard is only a guard for clients that choose to ask.
    """
    digest = hashlib.sha256(data).hexdigest()
    if expected_digest is not None and expected_digest != digest:
        raise AppError(
            409,
            "FILE_CHANGED",
            "That is not the file you just checked. Choose it again so the "
            "numbers you confirm describe what is imported.",
        )

    vendor = await _load_vendor(db, principal.company_id, vendor_id)

    # ── 1. read the sheet ────────────────────────────────────────────────
    raw_rows = list(_iter_rows(data, filename))
    if not raw_rows:
        raise _bad_request("The file is empty")

    header = {}
    for index, cell in enumerate(raw_rows[0]):
        key = next((c.key for c in COLUMNS if _norm(cell) in c.aliases), None)
        if key is not None and key not in header:
            header[key] = index
    missing = [
        _BY_KEY[k].header for k in REQUIRED_COLUMNS if k not in header
    ]
    if missing:
        found = ", ".join(_cell_text(c) for c in raw_rows[0] if _cell_text(c)) or "nothing"
        raise HTTPException(
            status_code=422,
            detail=(
                f"The sheet is missing {', '.join(missing)}. Found: {found}. "
                "Start from the template if you are not sure."
            ),
        )

    body_rows = raw_rows[1:]
    if len(body_rows) > MAX_TICKET_IMPORT_ROWS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"That file has {len(body_rows):,} rows. The importer takes "
                f"{MAX_TICKET_IMPORT_ROWS} at a time — split it and upload each "
                "part."
            ),
        )

    rows: list[_Row] = []
    for offset, raw in enumerate(body_rows):
        cells = {
            key: _cell_text(raw[i]) if i < len(raw) else ""
            for key, i in header.items()
        }
        if not any(cells.values()):
            continue  # A blank row is spacing, not a mistake.
        rows.append(_Row(number=offset + 2, cells=cells))

    # ── 2. everything decidable from the row alone ───────────────────────
    #
    # One instant for the whole file (hard rule 11), so two rows cannot be
    # judged against two different "todays" when an upload straddles midnight.
    today_ist = datetime.datetime.now(datetime.timezone.utc).astimezone(IST).date()
    for row in rows:
        _shape_check(row, today_ist)
        if row.reject is None:
            row.chain = parse_chain(row.cells.get("category", ""))
            if row.chain.code is not None:
                row.refuse(row.chain.code, row.chain.reason or "", "category")

    # ── 3. rows disagreeing with each other ──────────────────────────────
    _reconcile_chains(rows)

    # ── 4. the catalogue chains, in one call ─────────────────────────────
    live = [r for r in rows if r.reject is None and r.chain is not None]
    distinct: dict[tuple[str, ...], int] = {}
    chain_list: list[list[str]] = []
    for row in live:
        assert row.chain is not None
        if row.chain.key not in distinct:
            distinct[row.chain.key] = len(chain_list)
            chain_list.append(row.chain.segments)

    # On a dry run, and on a commit the caller has not confirmed, the function
    # is asked to resolve WITHOUT writing.
    resolve_dry = dry_run or not create_categories
    chain_results = await _resolve_chains(
        db, principal.company_id, principal.user_id, chain_list, dry_run=resolve_dry
    )

    for row in live:
        assert row.chain is not None
        result = chain_results[distinct[row.chain.key]]
        if result.conflict is not None:
            template = _CHAIN_CONFLICTS.get(
                result.conflict, "That category could not be used."
            )
            row.refuse(
                f"CATEGORY_{result.conflict}",
                template.format(
                    name=result.at,
                    child=result.child,
                    path=chain_path(row.chain.segments),
                    depth=result.depth,
                    length=len(result.at),
                ),
                "category",
            )
        elif result.node_id is None or (result.new_segments and not create_categories):
            row.refuse(
                "CATEGORY_NOT_FOUND",
                f'"{chain_path(row.chain.segments)}" is not in the catalogue '
                "yet. Confirm that it should be created, then import again.",
                "category",
            )
        else:
            row.node_id = result.node_id

    # ── 5. brands, products, serials, pincodes ───────────────────────────
    await _resolve_catalogue(db, principal, vendor, rows, submit_products, dry_run)

    # ── 6. rows this vendor has already imported ─────────────────────────
    await _mark_already_imported(db, principal.company_id, vendor, rows)

    # ── 7. what would happen ─────────────────────────────────────────────
    ready = [r for r in rows if r.reject is None and not r.already]
    quote = await quote_tickets(db, principal.company_id, count=len(ready))
    # The same question `intake_status` answers for the form: would the NEXT
    # ticket be refused? Derived from the quote rather than re-reading the
    # settings, so the two figures on this report cannot come from two reads.
    intake_paused = (
        quote.per_ticket > 0
        and quote.available - quote.per_ticket < -quote.floor
    )

    # The VENDOR's own line against this file. Summed from each row's stamped
    # price rather than a flat per-ticket figure, because that is what this side
    # of the money is: the ticket will owe what the model is priced at.
    #
    # Every `ready` row has a model by construction — a row naming a product the
    # catalogue does not hold yet, or one still awaiting pricing, is REFUSED
    # further up with "upload this row again once it is approved" — and an
    # approved model's `vendor_price_paise` is NOT NULL. So this total is exact,
    # not a floor.
    vendor_required = sum(
        r.model.vendor_price_paise for r in ready if r.model is not None
    )
    vendor_quote = await quote_vendor_credit(
        db, principal.company_id, vendor.id, required_paise=vendor_required
    )

    report = TicketImportReport(
        dryRun=dry_run,
        fileDigest=digest,
        rowsRead=len(rows),
        willImport=len(ready),
        alreadyImported=sum(1 for r in rows if r.already),
        rejected=sum(1 for r in rows if r.reject is not None),
        rejects=[
            r.reject
            for r in rows
            if r.reject is not None
        ][:MAX_TICKET_REJECTS_RETURNED],
        categoriesToCreate=_categories_to_create(live, distinct, chain_results),
        productsToSubmit=_products_to_submit(rows),
        creditsPerTicket=quote.per_ticket,
        creditsRequired=quote.required if show_credits else None,
        creditsAvailable=quote.available if show_credits else None,
        creditsShort=quote.short,
        intakePaused=intake_paused,
        # Both figures, to staff and vendor alike — see the schema for why this
        # side is not staff-only the way the company's is.
        vendorCreditRequiredPaise=vendor_required,
        vendorCreditAvailablePaise=vendor_quote.line.available_paise,
        vendorCreditShort=vendor_quote.short,
        vendorCreditPaused=vendor_quote.line.paused,
        duplicateRowsInFile=_count_repeats(rows),
    )

    if dry_run:
        # Nothing was written, but the ORM may hold objects a later read in this
        # session would otherwise see. Same guard `import_geography` ends on.
        db.expunge_all()
        return report

    await _commit(db, principal, vendor, rows, ready, report, submit_products)
    return report


@dataclasses.dataclass
class _ProductPlan:
    """A product the file names that the catalogue does not hold."""

    node_id: uuid.UUID
    brand_id: uuid.UUID
    brand_name: str
    name: str
    category_path: str
    service_types: set[str] = dataclasses.field(default_factory=set)
    rows: list[_Row] = dataclasses.field(default_factory=list)
    #: Filled on the commit, so the rows can say the product now exists.
    created_id: uuid.UUID | None = None


async def _resolve_catalogue(
    db: AsyncSession,
    principal: Principal,
    vendor: Vendor,
    rows: list[_Row],
    submit_products: bool,
    dry_run: bool,
) -> None:
    """Brand, product, serial and pincode — for the whole file, in five queries.

    The single-ticket path asks these per ticket, which is right for one and
    ruinous for five hundred. Every guard `_resolve_product` applies is still
    applied, in the same ORDER, and the two that need a row reuse its sentences
    verbatim:

      * the model must belong to THIS vendor — structural here, because the
        query filters `vendor_id`, so another vendor's product is simply absent
        and reads as `PRODUCT_NOT_FOUND`. That is deliberate: distinguishing
        "someone else's" from "does not exist" is the oracle
        `_resolve_product`'s comment warns about;
      * then approval, then the service type, then the serial — last, for the
        same reason it is last there.
    """
    from app.features.tickets.service import (
        pincode_refusal_text,
        serial_refusal_text,
    )

    live = [r for r in rows if r.reject is None and r.node_id is not None]

    # ── brands ───────────────────────────────────────────────────────────
    brands = list(
        await db.scalars(
            select(VendorBrand).where(
                VendorBrand.company_id == principal.company_id,
                VendorBrand.vendor_id == vendor.id,
                VendorBrand.deleted_at.is_(None),
            )
        )
    )
    by_name = {b.name.strip().lower(): b for b in brands}
    approved = [b for b in brands if b.approval_status == APPROVED]
    #: Omitting the brand is allowed only while there is exactly one approved —
    #: the same fallback `masters._resolve_brand` makes, for the same reason.
    sole = approved[0] if len(approved) == 1 else None

    for row in live:
        wanted = row.cells.get("brand", "").strip()
        if not wanted:
            if sole is None:
                row.refuse(
                    "BRAND_AMBIGUOUS",
                    "Name the brand. You have more than one approved brand, so "
                    "it cannot be guessed.",
                    "brand",
                )
                continue
            row.brand = sole
            continue
        found = by_name.get(wanted.lower())
        if found is None:
            row.refuse(
                "BRAND_NOT_FOUND",
                f'"{wanted}" is not one of {vendor.name}\'s approved brands. '
                "Add it on My brands, or use a brand you already have.",
                "brand",
            )
        elif found.approval_status == PENDING:
            row.refuse(
                "BRAND_NOT_APPROVED",
                f"{found.name} is still waiting for approval, so no product "
                "can carry it yet",
                "brand",
            )
        elif found.approval_status == REJECTED:
            row.refuse(
                "BRAND_REJECTED",
                f"{found.name} was not approved, so no product can carry it",
                "brand",
            )
        else:
            row.brand = found

    # ── products ─────────────────────────────────────────────────────────
    live = [r for r in rows if r.reject is None and r.node_id is not None]
    node_ids = {r.node_id for r in live if r.node_id is not None}
    models: dict[tuple[uuid.UUID, uuid.UUID, str], ProductModel] = {}
    if node_ids:
        for model in await db.scalars(
            select(ProductModel).where(
                ProductModel.company_id == principal.company_id,
                ProductModel.node_id.in_(node_ids),
                # THE pin. Another vendor's product is absent, not refused.
                ProductModel.vendor_id == vendor.id,
                ProductModel.deleted_at.is_(None),
            )
        ):
            models[(model.node_id, model.brand_id, model.name.strip().lower())] = model

    plans: dict[tuple[uuid.UUID, uuid.UUID, str], _ProductPlan] = {}
    for row in live:
        brand = row.brand
        if brand is None or row.node_id is None:
            continue
        name = row.cells.get("model", "").strip()
        if not name:
            row.refuse("MODEL_MISSING", "Name the product.", "model")
            continue
        key = (row.node_id, brand.id, name.lower())
        model = models.get(key)
        if model is None:
            assert row.chain is not None
            plan = plans.setdefault(
                key,
                _ProductPlan(
                    node_id=row.node_id,
                    brand_id=brand.id,
                    brand_name=brand.name,
                    name=name,
                    category_path=chain_path(row.chain.segments),
                ),
            )
            plan.service_types.add(row.cells.get("serviceType", "").strip())
            plan.rows.append(row)
            row.product_plan = plan
            if not submit_products:
                row.refuse(
                    "PRODUCT_NOT_FOUND",
                    f'"{name}" is not in the catalogue. Confirm that it should '
                    "be added, then import again.",
                    "model",
                )
            else:
                # It will be created by this import, and it will be PENDING, so
                # this row still cannot become a ticket. Two tenses, chosen by
                # `dry_run` — past tense in a dry run would be a lie.
                row.refuse(
                    "PRODUCT_AWAITING_APPROVAL",
                    (
                        f'"{name}" is not in the catalogue. It will be added '
                        "and will wait to be priced, so this row does not "
                        "import today."
                        if dry_run
                        else f'"{name}" has been added and is waiting to be '
                        "priced. Upload this row again once it is approved."
                    ),
                    "model",
                )
            continue

        if not model.is_active:
            row.refuse(
                "PRODUCT_INACTIVE",
                f"{model.name} is paused, so it cannot be ticketed.",
                "model",
            )
        elif model.approval_status == PENDING:
            row.refuse(
                "PRODUCT_AWAITING_APPROVAL",
                f"{model.name} is waiting to be priced, so it cannot be "
                "ticketed yet. Upload this row again once it is approved.",
                "model",
            )
        elif model.approval_status == REJECTED:
            row.refuse(
                "PRODUCT_REJECTED",
                f"{model.name} was not approved, so it cannot be ticketed. "
                "Fix it on My products and submit it again.",
                "model",
            )
        elif row.cells.get("serviceType", "").strip() not in (
            model.service_types or []
        ):
            row.refuse(
                "SERVICE_TYPE_UNSUPPORTED",
                f"{model.name} does not support "
                f'{row.cells.get("serviceType", "").strip()}. It supports '
                + ", ".join(model.service_types or [])
                + ".",
                "serviceType",
            )
        else:
            row.model_id = model.id
            row.model = model

    # ── serials: two queries for the file, not two per row ───────────────
    live = [r for r in rows if r.reject is None and r.model_id is not None]
    wanted_pairs = {
        (r.model_id, r.cells.get("serialNumber", "").strip().lower())
        for r in live
        if r.cells.get("serialNumber", "").strip()
    }
    hits: set[tuple[uuid.UUID, str]] = set()
    checked: set[uuid.UUID] = set()
    if wanted_pairs:
        model_ids = {m for m, _ in wanted_pairs}
        serials = {s for _, s in wanted_pairs}
        for model_id, serial in (
            await db.execute(
                select(
                    ProductModelSerial.product_model_id,
                    func.lower(ProductModelSerial.serial),
                ).where(
                    ProductModelSerial.company_id == principal.company_id,
                    ProductModelSerial.product_model_id.in_(model_ids),
                    func.lower(ProductModelSerial.serial).in_(serials),
                )
            )
        ).all():
            hits.add((model_id, serial))
        # Which models have ANY serials loaded — an empty list means unchecked,
        # not "nothing matches". One grouped query for the file.
        checked = {
            m
            for (m,) in (
                await db.execute(
                    select(ProductModelSerial.product_model_id)
                    .where(
                        ProductModelSerial.company_id == principal.company_id,
                        ProductModelSerial.product_model_id.in_(model_ids),
                    )
                    .group_by(ProductModelSerial.product_model_id)
                )
            ).all()
        }

    for row in live:
        serial = row.cells.get("serialNumber", "").strip()
        if not serial or row.model_id is None:
            continue
        if (row.model_id, serial.lower()) in hits:
            continue
        if row.model_id not in checked:
            continue  # Nothing loaded for this model: unchecked, so allowed.
        model = row.model
        row.refuse(
            "SERIAL_UNKNOWN",
            serial_refusal_text(serial, model.name if model else "this product"),
            "serialNumber",
        )

    # ── pincodes: one query for every distinct code ──────────────────────
    live = [r for r in rows if r.reject is None]
    codes = {r.cells.get("pincode", "").strip() for r in live}
    codes.discard("")
    known: set[str] = set()
    if codes:
        known = {
            c
            for (c,) in (
                await db.execute(
                    select(Pincode.code).where(
                        Pincode.code.in_(codes), Pincode.is_active.is_(True)
                    )
                )
            ).all()
        }
    for row in live:
        code = row.cells.get("pincode", "").strip()
        if code and code not in known:
            row.refuse("PINCODE_UNKNOWN", pincode_refusal_text(code), "pincode")


async def _mark_already_imported(
    db: AsyncSession, company_id: uuid.UUID, vendor: Vendor, rows: list[_Row]
) -> None:
    """Rows whose Reference is already on a ticket for this vendor.

    SKIPPED, not rejected, and the distinction is the whole point of the
    column: re-uploading the file after a product was approved is the normal
    way to use this importer, and every row that worked the first time comes
    back with it. A reject would read as a mistake; a skip reads as "already
    done", which is what it is. Neither raises a ticket, so neither is charged.
    """
    refs = {
        r.cells.get("externalRef", "").strip().lower()
        for r in rows
        if r.reject is None and r.cells.get("externalRef", "").strip()
    }
    if not refs:
        return
    seen = {
        ref
        for (ref,) in (
            await db.execute(
                select(func.lower(Ticket.external_ref)).where(
                    Ticket.company_id == company_id,
                    Ticket.vendor_id == vendor.id,
                    Ticket.deleted_at.is_(None),
                    func.lower(Ticket.external_ref).in_(refs),
                )
            )
        ).all()
    }
    in_file: set[str] = set()
    for row in rows:
        if row.reject is not None:
            continue
        ref = row.cells.get("externalRef", "").strip().lower()
        if not ref:
            continue
        if ref in seen:
            row.already = True
        elif ref in in_file:
            # Two rows in ONE file claiming the same reference. The partial
            # unique would refuse the second at the flush, so it is caught here
            # with a sentence instead of an IntegrityError.
            row.refuse(
                "DUPLICATE_REFERENCE_IN_FILE",
                f'Reference "{row.cells.get("externalRef", "").strip()}" is '
                "used by an earlier row. Each ticket needs its own.",
                "externalRef",
            )
        else:
            in_file.add(ref)


def _categories_to_create(
    live: list[_Row],
    distinct: dict[tuple[str, ...], int],
    results: list[_ChainResult],
) -> list[CategoryToCreate]:
    """The confirmation list: every chain with something new in it."""
    counts: dict[tuple[str, ...], int] = {}
    for row in live:
        assert row.chain is not None
        counts[row.chain.key] = counts.get(row.chain.key, 0) + 1

    out: list[CategoryToCreate] = []
    for row in live:
        assert row.chain is not None
        key = row.chain.key
        if any(c.path == chain_path(row.chain.segments) for c in out):
            continue
        result = results[distinct[key]]
        if result.conflict is not None:
            continue
        if not result.new_segments and not result.marked_leaf:
            continue
        # The level a technician certifies on. If THAT one is new, nobody
        # covers it and every ticket under it escalates on arrival.
        main = (
            row.chain.segments[CERTIFY_DEPTH]
            if len(row.chain.segments) > CERTIFY_DEPTH
            else None
        )
        out.append(
            CategoryToCreate(
                path=chain_path(row.chain.segments),
                newSegments=list(result.new_segments),
                depth=len(row.chain.segments),
                isLeaf=True,
                markedLeaf=result.marked_leaf,
                newMainSubcategory=main is not None and main in result.new_segments,
                rowCount=counts.get(key, 0),
            )
        )
    return out


def _products_to_submit(rows: list[_Row]) -> list[ProductToSubmit]:
    seen: dict[tuple, ProductToSubmit] = {}
    for row in rows:
        plan = row.product_plan
        if plan is None:
            continue
        key = (plan.node_id, plan.brand_id, plan.name.lower())
        if key not in seen:
            seen[key] = ProductToSubmit(
                categoryPath=plan.category_path,
                brandName=plan.brand_name,
                name=plan.name,
                serviceTypes=sorted(t for t in plan.service_types if t in SERVICE_TYPES),
                rowCount=len(plan.rows),
            )
    return list(seen.values())


async def _commit(
    db: AsyncSession,
    principal: Principal,
    vendor: Vendor,
    rows: list[_Row],
    ready: list[_Row],
    report: TicketImportReport,
    submit_products: bool,
) -> None:
    """Write the catalogue and the tickets, in ONE transaction.

    Categories, pending products and the good tickets commit together. A row
    rejected for "waiting to be priced" is a normal outcome and rolls nothing
    back; only a hard failure does, and then nothing at all is written. The
    chain creation is deliberately not committed separately — a file that
    created twelve categories and then failed to charge would leave a catalogue
    nobody asked for and no tickets to explain it.
    """
    from app.features.tickets.service import (
        push_pool_job,
        record_event,
    )

    now = datetime.datetime.now(datetime.timezone.utc)

    # ── the products nobody has priced yet ───────────────────────────────
    if submit_products:
        plans: dict[tuple, _ProductPlan] = {}
        for row in rows:
            if row.product_plan is not None:
                key = (
                    row.product_plan.node_id,
                    row.product_plan.brand_id,
                    row.product_plan.name.lower(),
                )
                plans.setdefault(key, row.product_plan)
        for plan in plans.values():
            types = sorted(t for t in plan.service_types if t in SERVICE_TYPES)
            db.add(
                ProductModel(
                    company_id=principal.company_id,
                    node_id=plan.node_id,
                    # THE pin, exactly as `submit_model` sets it: the vendor
                    # these tickets belong to, never a value from the file.
                    vendor_id=vendor.id,
                    brand_id=plan.brand_id,
                    name=plan.name,
                    service_types=types or list(SERVICE_TYPES[:1]),
                    # A vendor may never price a product, and `approved_is_priced`
                    # means an approved row must carry both figures — so this
                    # holds whoever ran the import.
                    technician_payout_paise=None,
                    vendor_price_paise=None,
                    approval_status=PENDING,
                    # `pending_was_submitted` requires it.
                    submitted_at=now,
                    created_by=principal.user_id,
                )
            )
        report.productsSubmitted = len(plans)

    report.categoriesCreated = sum(
        len(c.newSegments) for c in report.categoriesToCreate
    )

    if not ready:
        # Nothing to raise, but the catalogue writes above still belong to
        # somebody — commit them so a two-pass import makes progress on the
        # first pass rather than asking for a third.
        await db.commit()
        return

    # ── the tickets ──────────────────────────────────────────────────────
    #
    # One statement for every code in the file. `core.sequences.allocate`
    # increments the counter by N and returns the block; `tickets.service`'s own
    # note on `next_code` says the old COUNT(*) scheme raced and that "bulk
    # upload would have made that the normal case".
    codes = await allocate(db, principal.company_id, "ticket", len(ready))

    # Every node the file lands on, in ONE query — `node_path_ids` needs each
    # one's `ancestor_ids`, and asking per row is the per-row query this whole
    # module exists to avoid.
    nodes = {
        n.id: n
        for n in await db.scalars(
            select(ProductNode).where(
                ProductNode.company_id == principal.company_id,
                ProductNode.id.in_({r.node_id for r in ready if r.node_id}),
            )
        )
    }

    rules_by_node: dict[uuid.UUID, dict] = {}
    tickets: list[Ticket] = []
    for code, row in zip(codes, ready):
        body = row.body
        model = row.model
        assert body is not None and model is not None and row.node_id is not None
        node = nodes[row.node_id]
        if row.node_id not in rules_by_node:
            rules_by_node[row.node_id] = await resolve_rules(
                db, principal.company_id, row.node_id
            )
        ticket = Ticket(
            company_id=principal.company_id,
            code=code,
            vendor_id=vendor.id,
            node_id=node.id,
            model_id=model.id,
            node_path_ids=[*node.ancestor_ids, node.id],
            service_type=body.serviceType,
            description=body.description,
            serial_number=body.serialNumber,
            customer_name=body.customerName,
            customer_phone=body.customerPhone,
            address=body.address,
            city=body.city,
            state=body.state,
            pincode=body.pincode,
            # Never sent by this channel — see the module docstring.
            latitude=None,
            longitude=None,
            expected_date=body.expectedDate,
            service_level_hours=body.serviceLevelHours,
            sla_due_at=now + datetime.timedelta(hours=body.serviceLevelHours),
            # No slot column, so every imported ticket asks the customer.
            status="Slot Pending",
            slot_token=_token(),
            # The sweep sends it; this request cannot hold 500 WhatsApp round
            # trips. See `sweeps.sweep_pending_slot_requests`.
            slot_request_status="pending",
            technician_payout_paise=model.technician_payout_paise,
            vendor_price_paise=model.vendor_price_paise,
            rules_snapshot=rules_by_node[row.node_id],
            source="Excel",
            external_ref=row.cells.get("externalRef", "").strip() or None,
            created_by=principal.user_id,
        )
        tickets.append(ticket)
        db.add(ticket)

    # Autoflush is off (hard rule 10) and both the credit entries and the
    # events name the ticket id.
    await db.flush()

    # ── credits, all or nothing ──────────────────────────────────────────
    #
    # Taken LAST, and deliberately: `credits.lock` is `pg_advisory_xact_lock`,
    # held to the end of this transaction, and every single-ticket
    # `create_ticket` for this company takes the same key. Parsing, resolving
    # and building 500 rows first means a vendor raising one ticket by hand
    # waits milliseconds rather than the whole import.
    await charge_tickets(
        db,
        company_id=principal.company_id,
        ticket_ids=[t.id for t in tickets],
        by_user=principal.user_id,
    )

    # And the vendor's own line, whole-or-nothing for the same reason: a file is
    # one act, and half of it landing is worse than none of it.
    #
    # Nothing is written here. Every ticket above is flushed and non-terminal, so
    # the whole file is already inside the vendor's `reserved` and the check is
    # exact — `core.vendor_credits.assert_within_limit` explains the ordering.
    # Taken after the company's charge, so the two locks are always acquired in
    # the same order and two imports can never deadlock against each other.
    await assert_vendor_within_limit(
        db,
        company_id=principal.company_id,
        vendor_id=vendor.id,
        just_reserved_paise=sum(t.vendor_price_paise for t in tickets),
        ticket_count=len(tickets),
    )

    actor = vendor.name if principal.is_vendor else _staff_label(principal, vendor)
    for ticket in tickets:
        db.add(
            record_event(
                ticket,
                "created",
                actor_kind="vendor" if principal.is_vendor else "staff",
                actor_label=actor,
                note=(
                    f"{ticket.service_type} · {ticket.service_level_hours}h "
                    f"service level · {ticket.city} {ticket.pincode} · "
                    "imported from a spreadsheet"
                ),
                to_status=ticket.status,
                by_user=principal.user_id,
            )
        )

    # ── realtime, deduped ────────────────────────────────────────────────
    #
    # One pool frame per distinct (pincode, path): a technician's app needs to
    # know the pool moved, not how many times. Per-ticket console frames are
    # capped — 500 `pg_notify` calls would swamp every open console, and the
    # uploader is looking at this response rather than the board.
    buckets = {
        (t.pincode, tuple(t.node_path_ids)) for t in tickets
    }
    for pincode, path in buckets:
        await publish_pool_changed(
            db,
            company_id=principal.company_id,
            pincode=pincode,
            node_path_ids=list(path),
        )
    if len(tickets) <= MAX_PER_TICKET_FRAMES:
        for ticket in tickets:
            await publish_ticket_changed(db, ticket)

    await db.commit()

    # ── after the commit ─────────────────────────────────────────────────
    #
    # `push_pool_job` runs `technicians_covering` per ticket and sends per
    # technician, so it is called only for a small batch. Above the cap a
    # technician with the app OPEN still sees the jobs (the pool frames above),
    # and `sweep_unaccepted` still escalates anything nobody takes — but a
    # closed app is not told, and that is a real difference worth knowing. A
    # batched push ("12 new jobs near 500001") is the fix and is its own piece
    # of work.
    if len(tickets) <= MAX_PER_TICKET_FRAMES:
        for ticket in tickets:
            await push_pool_job(db, ticket)

    report.imported = len(tickets)
    report.slotRequestsQueued = len(tickets)


def _token() -> str:
    # 256 bits, the same as a technician invite, because it is the same kind of
    # secret: a URL somebody is trusted to hold.
    return secrets.token_urlsafe(32)


def _staff_label(principal: Principal, vendor: Vendor) -> str:
    name = getattr(principal.user, "full_name", None) or "Staff"
    return f"{name} (for {vendor.name})"


def _count_repeats(rows: list[_Row]) -> int:
    """Rows repeating another row's serial, product and customer.

    Reported, never rejected. `tickets.serial_number` is deliberately not
    unique — a second visit to one unit repeats it — so refusing a repeat would
    break the ordinary case to catch a copy-paste mistake. Surfacing it lets
    somebody look.
    """
    seen: set[tuple[str, str, str]] = set()
    repeats = 0
    for row in rows:
        key = (
            row.cells.get("serialNumber", "").strip().lower(),
            row.cells.get("model", "").strip().lower(),
            row.cells.get("customerPhone", "").strip(),
        )
        if not any(key):
            continue
        if key in seen:
            repeats += 1
        seen.add(key)
    return repeats
